"""Offline deterministic checks for Agent 2 Track 3."""

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from services.agent2.track3 import (
    Track3InputError,
    analysis_input_fingerprint,
    comparability_issues,
    compare_track3,
    export_track3_excel,
    load_synthetic_budget,
    parse_upload,
    review_readiness,
    role_mapping_table,
    synthetic_tenderers,
    validation_issues,
    validation_warnings,
    variance_severity,
)
from services.evidence import ROLES, load_records, map_role, rate_benchmark


def test_administrative_assistant_is_a_standard_comparison_role():
    assert "Administrative Assistant" in ROLES
    assert map_role("Administrative Assistant") == "Administrative Assistant"
    assert map_role("Admin Assistant") == "Administrative Assistant"
    assert map_role("Admin Executive") == "Administrative Assistant"


def test_track3_catalogue_covers_track4_building_roles():
    records = load_records()
    frame = pd.DataFrame(
        [{"Role": role} for role in [
            "Senior Facilities Manager", "BMS/BOS Officer", "Cleaning Supervisor",
            "Helpdesk Officer", "M&E Engineer", "Property Executive", "Security Supervisor",
        ]]
    )
    mapping = role_mapping_table({"Tenderer A": frame})
    for role in [
        "Senior Facilities Manager", "BMS/BOS Officer", "Cleaning Supervisor",
        "Helpdesk Officer", "M&E Engineer", "Property Executive", "Security Supervisor",
    ]:
        row = mapping[mapping["Role title in submission"] == role].iloc[0]
        assert row["Compare as"] == role
        assert row["Why this is suggested"] == "Title already matches the standard role list"
        assert rate_benchmark(records, role, year=2026)["available"]


def test_validation_warnings_accept_and_explain_unusual_inputs():
    frame = pd.DataFrame(
        [
            {
                "Role": "Senior Facilities Manager",
                "Contract Year": 1,
                "Quantity": 1,
                "Quantity Unit": "headcount",
                "Monthly Rate (SGD)": 18500,
                "Source": "officer schedule",
                "Officer Justification": "Critical-site coverage confirmed.",
                "Officer Assumption Confirmed": "Yes",
            },
            {
                "Role": "Administrative Assistant",
                "Contract Year": 1,
                "Quantity": 0.5,
                "Quantity Unit": "fte",
                "Monthly Rate (SGD)": 3900,
                "Source": "officer schedule",
                "Officer Justification": "Shared deployment.",
                "Officer Assumption Confirmed": "Yes",
            },
            {
                "Role": "Drone Façade Inspection Specialist",
                "Contract Year": 1,
                "Quantity": 1,
                "Quantity Unit": "headcount",
                "Monthly Rate (SGD)": 12800,
                "Source": "officer schedule",
                "Officer Justification": "Specialist requirement; no comparable benchmark is available.",
                "Officer Assumption Confirmed": "Yes",
            },
        ]
    )

    assert validation_issues(frame).empty
    warnings = validation_warnings(frame)
    assert len(warnings) == 3
    assert warnings["Review note"].str.contains("High monthly rate").any()
    assert warnings["Review note"].str.contains("0.5 FTE").any()
    assert warnings["Review note"].str.contains("no benchmark will be invented").any()


def _ready_context():
    budget, _ = load_synthetic_budget()
    tenderers = synthetic_tenderers(budget)
    mappings = role_mapping_table({"Approved budget": budget, **tenderers})
    project = {
        "name": "Synthetic Track 3",
        "duration": 6,
        "base_years": 4,
        "option_years": 2,
        "start": "2027-01-01",
        "currency": "SGD",
        "budget_approval_reference": "SYNTHETIC",
        "budget_revision": "1",
    }
    return project, budget, tenderers, mappings


def test_review_readiness_is_deterministic_and_marks_changed_inputs_stale():
    project, budget, tenderers, mappings = _ready_context()
    result = compare_track3(budget, tenderers)
    fingerprint = analysis_input_fingerprint(project, budget, tenderers, mappings)

    current = review_readiness(
        project, budget, tenderers, mappings,
        budget_confirmed=True, mappings_confirmed=True,
        result=result, analysed_fingerprint=fingerprint,
    )
    assert current.analysis_is_current
    assert current.status == "review_in_progress"

    changed = tenderers.copy()
    changed["Tenderer A"] = changed["Tenderer A"].copy()
    changed["Tenderer A"].loc[0, "Quantity"] += 1
    stale = review_readiness(
        project, budget, changed, mappings,
        budget_confirmed=True, mappings_confirmed=True,
        result=result, analysed_fingerprint=fingerprint,
    )
    assert stale.status == "stale"
    assert not stale.analysis_is_current


def test_review_readiness_explains_missing_prerequisites():
    readiness = review_readiness(
        {}, None, {}, None,
        budget_confirmed=False, mappings_confirmed=False,
        result=None, analysed_fingerprint=None,
    )
    assert readiness.status == "checks_required"
    assert "Create the review project." in readiness.reasons
    assert "Add the approved manpower budget." in readiness.reasons


def test_synthetic_track3_compares_multiple_tenderers_and_flags_risks():
    budget, _ = load_synthetic_budget()
    result = compare_track3(budget, synthetic_tenderers(budget))
    assert set(result.overall["Tenderer"]) == {"Tenderer A", "Tenderer B", "Tenderer C"}
    assert "⚠ Below budget — clarify assumptions" in set(result.overall["Assessment"])
    assert "⚠ Above approved budget — review required" in set(result.overall["Assessment"])
    assert not result.findings.empty
    assert not result.questions.empty
    assert any("Technician" in value for value in result.findings["Finding"])
    assert any("contract year 6" in value for value in result.findings["Finding"])


def test_missing_values_are_flagged_and_not_changed_to_zero():
    frame = pd.DataFrame(
        [{"Role": "Technician", "Contract Year": 1, "Quantity": None, "Quantity Unit": "headcount", "Monthly Rate (SGD)": None, "Source": "test.csv"}]
    )
    issues = validation_issues(frame)
    assert set(issues["Field"]) == {"Quantity", "Monthly Rate (SGD)"}
    assert pd.isna(frame.loc[0, "Quantity"])


def test_fractional_headcount_requires_explicit_fte_but_half_fte_is_accepted():
    base = {
        "Role": "Facilities Manager",
        "Contract Year": 1,
        "Quantity": 0.5,
        "Monthly Rate (SGD)": 8200,
        "Source": "validation",
    }
    ambiguous = pd.DataFrame([{**base, "Quantity Unit": "headcount"}])
    assert "Fractional quantity" in " ".join(validation_issues(ambiguous)["Issue"])
    explicit = pd.DataFrame([{**base, "Quantity Unit": "fte"}])
    assert validation_issues(explicit).empty


def test_contract_duration_coverage_and_double_escalation_are_flagged():
    incomplete = pd.DataFrame(
        [
            {
                "Role": "Technician", "Contract Year": year, "Quantity": 5,
                "Quantity Unit": "headcount", "Monthly Rate (SGD)": 4100,
                "Source": "validation", "Contract Duration (Years)": 3,
            }
            for year in (1, 2)
        ]
    )
    assert "missing annual rate(s)" in " ".join(validation_issues(incomplete)["Issue"])
    double_escalated = incomplete.iloc[[0]].copy()
    double_escalated["Contract Duration (Years)"] = 1
    double_escalated["Quoted Pricing Year"] = 2099
    double_escalated["Additional Escalation Applied (%)"] = 0.05
    assert "not applied twice" in " ".join(validation_issues(double_escalated)["Issue"])


def test_csv_upload_aliases_and_excel_export():
    content = b"Manpower Role,Year,Headcount,Monthly Rate (SGD)\nTechnician,1,4,4200\n"
    frame = parse_upload("Tenderer A.csv", content)
    assert list(frame.columns) == ["Role", "Contract Year", "Quantity", "Quantity Unit", "Monthly Rate (SGD)", "Source"]
    assert frame.loc[0, "Role"] == "Technician"

    budget, _ = load_synthetic_budget()
    result = compare_track3(budget, synthetic_tenderers(budget))
    binary = export_track3_excel({"name": "Synthetic Track 3"}, result, ["Synthetic only"])
    workbook = load_workbook(BytesIO(binary), read_only=True, data_only=False)
    assert {"Project Overview", "Overall Comparison", "Role Comparison", "Findings", "Clarification Questions", "Limitations"} <= set(workbook.sheetnames)


def test_excel_template_header_can_follow_title_and_instruction_rows():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Year 1"
    sheet.append(["Contract Year 1 Manpower Schedule"])
    sheet.append(["Complete the yellow cells."])
    sheet.append([])
    sheet.append([])
    sheet.append(["Role", "Contract Year", "Quantity", "Quantity Unit", "Monthly Rate (SGD)", "Source"])
    sheet.append(["Technician", 1, 4, "headcount", 4200, "Year 1 row 1"])
    output = BytesIO()
    workbook.save(output)
    frame = parse_upload("completed_template.xlsx", output.getvalue())
    assert len(frame) == 1
    assert frame.loc[0, "Role"] == "Technician"


def test_missing_contract_year_is_blocked_before_cost_comparison():
    budget, _ = load_synthetic_budget()
    tender = synthetic_tenderers(budget)["Tenderer A"]
    tender = tender[tender["Contract Year"] < tender["Contract Year"].max()]
    with pytest.raises(Track3InputError, match="comparability issue"):
        compare_track3(budget, {"Tenderer A": tender})


def test_duplicate_role_year_is_blocked_before_cost_comparison():
    budget, _ = load_synthetic_budget()
    tender = synthetic_tenderers(budget)["Tenderer A"]
    tender = pd.concat([tender, tender.iloc[[0]]], ignore_index=True)
    issues = comparability_issues(budget, {"Tenderer A": tender})
    assert "Calculation integrity" in set(issues["Category"])
    with pytest.raises(Track3InputError, match="Duplicate rows"):
        compare_track3(budget, {"Tenderer A": tender})


def test_headcount_fte_mismatch_is_blocked():
    budget, _ = load_synthetic_budget()
    tender = synthetic_tenderers(budget)["Tenderer A"].copy()
    tender["Quantity Unit"] = "fte"
    with pytest.raises(Track3InputError, match="Quantity unit differs"):
        compare_track3(budget, {"Tenderer A": tender})


def test_project_duration_must_match_schedule_coverage():
    budget, _ = load_synthetic_budget()
    tender = synthetic_tenderers(budget)["Tenderer A"]
    with pytest.raises(Track3InputError, match="missing year"):
        compare_track3(budget, {"Tenderer A": tender}, expected_years=7)


def test_traceable_export_includes_verified_inputs_mapping_assumptions_and_audit():
    budget, _ = load_synthetic_budget()
    tenderers = synthetic_tenderers(budget)
    result = compare_track3(budget, tenderers, benchmark_year=2027)
    mappings = pd.DataFrame(
        [{"Original Role": "Technician", "Standard Role": "Technician", "Confirmed": True}]
    )
    binary = export_track3_excel(
        {
            "name": "Synthetic Track 3",
            "budget_confirmed": True,
            "budget_approval_reference": "SYNTHETIC",
            "budget_revision": 1,
        },
        result,
        ["Synthetic only"],
        budget=budget,
        tenderers=tenderers,
        role_mappings=mappings,
    )
    workbook = load_workbook(BytesIO(binary), read_only=True, data_only=False)
    assert {
        "Verified Budget",
        "Verified Tender Data",
        "Role Mapping",
        "Assumptions",
        "Audit Trail",
    } <= set(workbook.sheetnames)


def test_material_findings_are_ranked_and_questions_include_figures():
    budget, _ = load_synthetic_budget()
    result = compare_track3(budget, synthetic_tenderers(budget))
    assert {"Cost Variance to Budget (SGD)", "Cost Variance to Budget (%)", "Source Reference"} <= set(result.findings.columns)
    assert result.findings.iloc[0]["Cost Variance to Budget (SGD)"] == result.findings["Cost Variance to Budget (SGD)"].max()
    assert any("SGD" in question for question in result.questions["Clarification Question"])
    assert {"S/N", "Priority"} <= set(result.questions.columns)


def test_variance_severity_uses_percentage_and_dollar_impact():
    assert variance_severity(411_775.32, 0.164) == "High"
    assert variance_severity(75_159.84, 0.03) == "Medium"
    assert variance_severity(10_000, 0.03) == "Low"


def test_role_matching_keeps_tenderer_and_budget_titles_separate():
    budget, _ = load_synthetic_budget()
    tenderers = synthetic_tenderers(budget)
    mapping = role_mapping_table({"Approved budget": budget, **tenderers})
    assert {"Submitted by", "Role title in submission", "Compare as", "Why this is suggested"} == set(
        mapping.columns
    )
    assert set(mapping["Submitted by"]) == {"Approved budget", "Tenderer A", "Tenderer B", "Tenderer C"}
    assert not mapping.duplicated(["Submitted by", "Role title in submission"]).any()


@pytest.mark.parametrize(
    ("extra_columns", "message"),
    [
        ("Currency,Pricing Frequency\nUSD,monthly", "only SGD"),
        ("Currency,Pricing Frequency\nSGD,annual", "rates must be monthly"),
    ],
)
def test_upload_blocks_incomparable_currency_and_pricing_frequency(extra_columns, message):
    headings, values = extra_columns.splitlines()
    content = (
        f"Role,Contract Year,Quantity,Quantity Unit,Monthly Rate (SGD),{headings}\n"
        f"Technician,1,4,headcount,4200,{values}\n"
    ).encode()
    with pytest.raises(Track3InputError, match=message):
        parse_upload("Tenderer A.csv", content)


def test_upload_blocks_annual_rate_heading():
    content = b"Role,Contract Year,Quantity,Quantity Unit,Annual Rate (SGD)\nTechnician,1,4,headcount,50400\n"
    with pytest.raises(Track3InputError, match="annual pricing column"):
        parse_upload("Tenderer A.csv", content)


def test_upload_does_not_silently_assume_currency():
    content = b"Role,Contract Year,Quantity,Quantity Unit,Monthly Rate\nTechnician,1,4,headcount,4200\n"
    with pytest.raises(Track3InputError, match="currency is not established"):
        parse_upload("Tenderer A.csv", content)


def test_top_level_json_array_is_supported_when_pricing_basis_is_explicit():
    content = (
        b'[{"Role":"Technician","Contract Year":1,"Quantity":4,'
        b'"Quantity Unit":"headcount","Monthly Rate (SGD)":4200}]'
    )
    frame = parse_upload("Tenderer A.json", content)
    assert frame.loc[0, "Monthly Rate (SGD)"] == 4200


def test_excel_upload_combines_only_schedule_worksheets_with_sheet_provenance():
    output = BytesIO()
    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        pd.DataFrame({"Note": ["Instructions only"]}).to_excel(writer, sheet_name="Read me", index=False)
        pd.DataFrame(
            [{"Role": "Technician", "Contract Year": 1, "Quantity": 4,
              "Quantity Unit": "headcount", "Monthly Rate (SGD)": 4200}]
        ).to_excel(writer, sheet_name="Base", index=False)
        pd.DataFrame(
            [{"Role": "Technician", "Contract Year": 2, "Quantity": 4,
              "Quantity Unit": "headcount", "Monthly Rate (SGD)": 4300}]
        ).to_excel(writer, sheet_name="Option", index=False)

    frame = parse_upload("Tenderer A.xlsx", output.getvalue())
    assert list(frame["Contract Year"]) == [1, 2]
    assert "sheet 'Base'" in frame.loc[0, "Source"]
    assert "sheet 'Option'" in frame.loc[1, "Source"]


def test_bct_composition_tabs_extract_year_roles_and_weighted_site_rates():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "5. Yr1 Composition Rates"
    sheet["A1"] = "DETAILED COST: COMPOSITION OF RATES ON BREAKDOWN COST OF TENDER"
    sheet["A2"] = "INITIAL TERM - YEAR 1"
    sheet["C3"] = "Description"
    sheet["E3"] = "Site A"
    sheet["H3"] = "Site B"
    sheet["D4"] = "Unit"
    sheet["E4"] = "Unit Rate per month (A1)"
    sheet["F4"] = "Quantity (B1)"
    sheet["G4"] = "Total Cost Per Month"
    sheet["H4"] = "Unit Rate per month (A2)"
    sheet["I4"] = "Quantity (B2)"
    sheet["J4"] = "Total Cost Per Month"
    sheet.append([])
    sheet["C6"] = "a) Facility Manager"
    sheet["D6"] = "pax"
    sheet["E6"] = 9000
    sheet["F6"] = 0.25
    sheet["H6"] = 10000
    sheet["I6"] = 0.75
    sheet["C7"] = "b) Technician"
    sheet["D7"] = "pax"
    sheet["E7"] = 4000
    sheet["F7"] = 2
    sheet["C8"] = "c) Training"
    sheet["D8"] = "lump sum"
    sheet["E8"] = 2500
    sheet["F8"] = 1
    output = BytesIO()
    workbook.save(output)

    frame = parse_upload("BCT Tenderer.xlsx", output.getvalue())

    assert list(frame["Role"]) == ["Facility Manager", "Technician"]
    assert set(frame["Contract Year"]) == {1}
    facility = frame.loc[frame["Role"] == "Facility Manager"].iloc[0]
    assert facility["Quantity"] == pytest.approx(1.0)
    assert facility["Monthly Rate (SGD)"] == pytest.approx(9750.0)
    assert facility["Quantity Unit"] == "headcount"
    assert "Site A" in facility["Source"] and "Site B" in facility["Source"]


def test_bct_composition_workbook_combines_year_tabs():
    workbook = Workbook()
    first = workbook.active
    first.title = "5. Yr1 Composition Rates"
    second = workbook.create_sheet("6. Yr2 Composition Rates")
    for year, sheet in ((1, first), (2, second)):
        sheet["A1"] = f"INITIAL TERM - YEAR {year}"
        sheet["C3"] = "Description"
        sheet["E3"] = "Site A"
        sheet["D4"] = "Unit"
        sheet["E4"] = "Unit Rate per month (A1)"
        sheet["F4"] = "Quantity (B1)"
        sheet["G4"] = "Total Cost Per Month"
        sheet["C6"] = "Technician"
        sheet["D6"] = "pax"
        sheet["E6"] = 4000 + year * 100
        sheet["F6"] = 2
    output = BytesIO()
    workbook.save(output)

    frame = parse_upload("multi-year.xlsx", output.getvalue())

    assert list(frame["Contract Year"]) == [1, 2]
    assert list(frame["Monthly Rate (SGD)"]) == [4100, 4200]


def test_composition_parser_infers_role_column_and_year_when_labels_and_tab_differ():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Commercial Pricing"
    sheet["A2"] = "OPTION PERIOD — YEAR 3"
    sheet["E3"] = "North Site"
    sheet["D4"] = "UOM"
    sheet["E4"] = "Monthly unit price"
    sheet["F4"] = "No. of pax"
    sheet["G4"] = "Extended monthly amount"
    sheet["C6"] = "01. Property Executive"
    sheet["D6"] = "FTE"
    sheet["E6"] = 6200
    sheet["F6"] = 1.5
    output = BytesIO()
    workbook.save(output)

    frame = parse_upload("alternate-layout.xlsx", output.getvalue())

    assert len(frame) == 1
    assert frame.loc[0, "Role"] == "Property Executive"
    assert frame.loc[0, "Contract Year"] == 3
    assert frame.loc[0, "Quantity"] == pytest.approx(1.5)
    assert frame.loc[0, "Monthly Rate (SGD)"] == pytest.approx(6200)


def test_detected_composition_table_without_year_requires_confirmation():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Commercial Pricing"
    sheet["C3"] = "Position"
    sheet["D4"] = "Unit"
    sheet["E4"] = "Rate per month"
    sheet["F4"] = "Headcount"
    sheet["C6"] = "Technician"
    sheet["D6"] = "headcount"
    sheet["E6"] = 4000
    sheet["F6"] = 2
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(Track3InputError, match="contract year could not be established"):
        parse_upload("missing-year.xlsx", output.getvalue())


def test_corrupt_excel_has_user_facing_error():
    with pytest.raises(Track3InputError, match="could not be read"):
        parse_upload("broken.xlsx", b"not an excel workbook")
