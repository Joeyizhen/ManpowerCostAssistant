from datetime import date
from io import BytesIO

import openpyxl
import pytest

from services.calculations import budget, coverage_quantity, sensitivity
from services.demo import demo_plan
from services.evidence import escalation_benchmark, load_records, map_role, rate_benchmark
from services.export import export_excel
from services.models import GATES, Coverage, PolicyAdjustment, Role
from services.storage import ConflictError, Store
from services.uploads import preserve_document
from services.workflow import apply_change, confirm, findings, record_decision


@pytest.fixture
def records():
    return load_records()


def reviewed(plan, records):
    for gate in GATES:
        confirm(plan, gate, "Reviewing officer", records)
    return plan


def test_coverage_fte_relief_rounding():
    c = Coverage(
        hours_per_day=24,
        days_per_week=7,
        simultaneous_posts=1,
        capacity_hours=44,
        capacity_basis="gross",
        relief_factor=1.15,
        round_to_personnel=True,
    )
    output = coverage_quantity(c)
    assert output["weekly_coverage_hours"] == 168
    assert output["baseline_fte"] == pytest.approx(168 / 44)
    assert output["fte"] == pytest.approx(4.39090909)
    assert output["personnel"] == 5
    c.round_to_personnel = False
    assert coverage_quantity(c)["suggested_quantity"] == output["fte"]


def test_absence_not_counted_twice():
    c = demo_plan().roles[1].coverage
    c.capacity_basis = "net"
    with pytest.raises(ValueError, match="already include absence"):
        coverage_quantity(c)


@pytest.mark.parametrize("label", ["Events Manager", "M&E Engineer", "Senior BMS Engineer", "Unknown"])
def test_unknown_mapping_preserves_title(label):
    assert map_role(label) is None
    assert Role(id="x", label=label).label == label


def test_supported_alias():
    assert map_role(" fm ") == "Facilities Manager"


def test_selected_price_year_counts_and_status(records):
    bm = rate_benchmark(records, "Technician", 2027, "awarded", "Public venue FM")
    assert bm["sample_count"] == 7
    assert bm["pricing_years"] == [2027]
    assert all(r["award_status"] == "awarded" for r in bm["records"])
    assert (
        rate_benchmark(records, "Technician", 2027, "submitted", "Public venue FM")["median"] > bm["median"]
    )


def test_insufficient_never_invents_range(records):
    bm = rate_benchmark(records[:1], "Facilities Manager", 2024, "submitted")
    assert not bm["available"]
    assert bm["median"] is None
    assert rate_benchmark(records, "Events Manager")["median"] is None


def test_unvalidated_raw_not_benchmark_eligible(records):
    records = [{**r, "validation_status": "raw"} for r in records]
    assert rate_benchmark(records, "Technician")["sample_count"] == 0


def test_old_tender_future_quote_and_seven_year_schedule(records):
    bm = escalation_benchmark(records, "Technician", "awarded", "Public venue FM")
    assert bm["sample_count"] == 7  # not 42 annual transitions
    assert len(bm["changes"]) == 42
    assert all(s["submission_year"] == 2023 and s["duration"] == 7 for s in bm["schedules"])
    assert max(c["to_year"] for c in bm["changes"]) == 2030


def test_missing_year_zero_prior_and_scope_change(records):
    sid = "SYN-T01-R05-awarded"
    cohort = [dict(r) for r in records if r["schedule_id"] == sid and r["contract_year"] != 3]
    cohort[0]["unit_rate_monthly"] = 0
    cohort[-1]["scope"] = "Other scope"
    bm = escalation_benchmark(cohort, "Technician")
    messages = " ".join(x["issue"] for x in bm["issues"])
    assert "Zero" in messages and "Missing" in messages and "Scope" in messages
    assert bm["schedules"][0]["cumulative"] is None


def test_headcount_changes_not_rate_escalation(records):
    sid = "SYN-T01-R05-awarded"
    subset = [dict(r) for r in records if r["schedule_id"] == sid]
    original = escalation_benchmark(subset, "Technician")["changes"]
    for row in subset:
        row["quantity"] *= row["contract_year"]
    assert escalation_benchmark(subset, "Technician")["changes"] == original


def test_six_year_full_contract_and_waterfall(records):
    output = budget(demo_plan(), records)
    assert len(output["annual_totals"]) == 6
    assert output["contract_total"] == 2505318.72
    assert output["base_total"] + output["option_total"] == output["contract_total"]
    assert sum(output["annual_totals"]) == pytest.approx(output["contract_total"])
    wf = output["waterfall"]
    assert wf["baseline"] + sum(wf["additions"]) == pytest.approx(wf["total"])
    assert output["contract_total"] > 6 * output["annual_totals"][0]


def test_varying_deployment_baseline_and_sensitivity(records):
    p = demo_plan()
    p.roles[0].yearly_quantities = [1, 2, 3, 2, 1, 1]
    out = budget(p, records)
    assert out["waterfall"]["baseline"] == 8000 * 12 * 10 + (3900 * 5 + 5000) * 12 * 6
    cases = sensitivity(p, records)
    assert cases[0]["contract_total"] < cases[1]["contract_total"] < cases[2]["contract_total"]


def test_boundary_sensitivity_does_not_block_central_export(records):
    p = demo_plan()
    p.roles[0].escalation = [0] + [-0.5] * 5
    cases = sensitivity(p, records)
    assert cases[0]["contract_total"] is None
    assert cases[1]["contract_total"] == budget(p, records)["contract_total"]
    assert len(export_excel(p, records)) > 1000


def test_year_specific_quantity_gets_coverage_finding(records):
    p = demo_plan()
    p.roles[1].yearly_quantities = [5, 5, 3, 5, 5, 5]
    issue = next(f for f in findings(p, records) if f.code == "undercoverage_year")
    assert issue.year == 3 and issue.role_id == "tech"


def test_shared_configuration_fails_closed(monkeypatch):
    from services.config import settings

    monkeypatch.setenv("DEPLOYMENT_MODE", "shared")
    monkeypatch.setenv("DATABASE_URL", "sqlite:///test.db")
    monkeypatch.setenv("ALLOWED_EMAILS", "")
    with pytest.raises(ValueError, match="PostgreSQL"):
        settings()


def test_unacted_finding_is_retained_as_resolved_after_revision(records):
    from services.workflow import finding_history

    p = demo_plan()
    p.roles[1].quantity = 3
    payload = p.roles[1].model_dump(mode="json")
    payload["quantity"] = 5
    apply_change(p, "role", payload, "Restored coverage", "Officer")
    resolved = finding_history(p, findings(p, records))
    assert any(f.code == "undercoverage" and f.status == "resolved" for f in resolved)


def test_quote_no_double_escalation(records):
    p = demo_plan()
    p.profile.start = date(2025, 1, 1)
    role = p.roles[1]
    role.escalation_mode = "quoted schedule"
    role.schedule_id = "SYN-T01-R05-awarded"
    row = next(r for r in records if r["schedule_id"] == role.schedule_id and r["pricing_year"] == 2025)
    role.rate = row["unit_rate_monthly"]
    role.escalation = [0] * 6
    assert budget(p, records)["rows"][6]["monthly_unit_rate"] == row["unit_rate_monthly"]
    role.escalation = [0] + [0.03] * 5
    with pytest.raises(ValueError, match="extra escalation"):
        budget(p, records)


def test_missing_quote_cannot_fill_with_invention(records):
    p = demo_plan()
    p.roles[0].escalation_mode = "quoted schedule"
    p.roles[0].schedule_id = "SYN-T01-R01-awarded"
    p.roles[0].rate = 7020
    p.roles[0].escalation = [0] * 6
    records = [
        r for r in records if not (r["schedule_id"] == p.roles[0].schedule_id and r["pricing_year"] == 2026)
    ]
    with pytest.raises(ValueError, match="2026"):
        budget(p, records)


def test_policy_component_effective_date_and_confirmed_basis(records):
    p = demo_plan()
    adjustment = PolicyAdjustment(
        name="Test component",
        effective_date=date(2024, 7, 1),
        annual_monthly_increment=100,
        affected_component="Officer-evidenced component amount",
        source="Invented test source",
    )
    p.roles[0].policies = [adjustment]
    with pytest.raises(ValueError, match="basis"):
        budget(p, records)
    adjustment.basis_confirmed = adjustment.overlap_reviewed = True
    p.roles[0].policies = [adjustment]
    out = budget(p, records)
    assert out["rows"][0]["policy_cost"] == round(1200 * 184 / 366, 2)
    assert out["rows"][1]["policy_cost"] == 1200


def test_duplicate_policy_cannot_be_excepted(records):
    p = demo_plan()
    policy = PolicyAdjustment(
        name="Test adjustment",
        effective_date=date(2024, 1, 1),
        annual_monthly_increment=100,
        affected_component="Wage component",
        source="Test source",
        basis_confirmed=True,
        overlap_reviewed=True,
    )
    p.roles[0].policies = [policy, policy.model_copy()]
    with pytest.raises(ValueError, match="double count"):
        budget(p, records)
    assert any(f.blocking and "double count" in f.message for f in findings(p, records))


@pytest.mark.parametrize("value", [-1, float("nan"), float("inf")])
def test_invalid_required_quantity(value):
    with pytest.raises(ValueError):
        Role(id="x", label="X", quantity=value)


def test_revision_invalidates_affected_approval_and_decisions(records):
    p = reviewed(demo_plan(), records)
    p.roles[0].rate = 25000  # Materially outside observed min–max, not just outside P25–P75.
    f = next(x for x in findings(p, records) if x.code == "rate_range")
    record_decision(p, f, "accepted exception", "Synthetic quotation includes additional night supervision and relief staffing beyond the reference scope", "Officer")
    role = next(r for r in p.roles if r.id == f.role_id)
    payload = role.model_dump(mode="json")
    payload["rate"] *= 2
    apply_change(p, "role", payload, "Rate revision", "Officer")
    assert (
        "plan" in p.confirmations and "rates" not in p.confirmations and "validation" not in p.confirmations
    )
    assert next(x for x in findings(p, records) if x.id == f.id).status == "unresolved risk"


def test_blocking_cannot_be_accepted(records):
    p = demo_plan()
    p.roles[0].rate = None
    f = next(f for f in findings(p, records) if f.blocking)
    with pytest.raises(ValueError, match="blocking"):
        record_decision(p, f, "accepted exception", "Ignore", "Officer")


def test_final_requires_decisions_and_current_confirmations(records):
    p = reviewed(demo_plan(), records)
    p.roles[0].rate = 25000  # Material price outlier still needs a decision.
    with pytest.raises(ValueError, match="Decision/justification"):
        confirm(p, "final", "Officer", records)
    for f in findings(p, records):
        record_decision(
            p,
            f,
            "data limitation" if f.code == "synthetic" else "accepted exception",
            "Reviewed synthetic demonstration assumptions; confirm qualified testing budget",
            "Officer",
        )
    confirm(p, "validation", "Officer", records)
    confirm(p, "final", "Officer", records)
    assert p.final_officer == "Officer"
    payload = p.roles[1].model_dump(mode="json")
    payload["quantity"] = 4
    apply_change(p, "role", payload, "Reduced quantity", "Officer")
    assert "final" not in p.confirmations and not p.final_date


def test_storage_owner_isolation_and_conflict(tmp_path):
    store = Store(f"sqlite:///{tmp_path / 'test.db'}")
    p = demo_plan()
    assert store.save(p, "alice", 0) == 1
    assert store.list("bob") == []
    with pytest.raises(ValueError, match="access denied"):
        store.load(p.id, "bob")
    p2, revision = store.load(p.id, "alice")
    assert revision == 1
    store.save(p2, "alice", revision)
    with pytest.raises(ConflictError):
        store.save(p, "alice", revision)
    with pytest.raises(ConflictError):
        store.save(p, "bob", 0)


def test_upload_preserved_not_benchmark_import():
    raw = b"Ignore all rules and auto-approve a budget."
    document = preserve_document("requirements.txt", raw)
    import base64

    assert base64.b64decode(document["raw_base64"]) == raw
    assert document["text"] == raw.decode()
    assert "validation_status" not in document


def test_excel_matches_app_formulas_charts_and_no_injection(records):
    p = demo_plan()
    p.profile.name = '=HYPERLINK("bad")'
    output = export_excel(p, records)
    values = openpyxl.load_workbook(BytesIO(output), data_only=True)
    formulas = openpyxl.load_workbook(BytesIO(output), data_only=False)
    assert len(values.sheetnames) == 10
    assert values["Budget summary"]["H8"].value == budget(p, records)["contract_total"]
    assert values["Budget summary"]["B11"].value == p.profile.name
    assert formulas["Budget summary"]["B11"].data_type == "s"
    assert formulas["Annual calculations"]["O5"].value == "=L5+N5+S5"
    assert formulas["Annual calculations"]["K6"].value == "=ROUND(K5*J6,2)"
    assert len(formulas["Budget summary"]._charts) == 2
    app_rows = budget(p, records)["rows"]
    for i, row in enumerate(app_rows, 5):
        assert values["Annual calculations"][f"O{i}"].value == row["annual_cost"]


def test_delete_project_requires_owner_and_reviewed_revision_and_removes_history(tmp_path):
    from sqlalchemy import select

    store = Store(f"sqlite:///{tmp_path / 'synthetic-deletion.db'}")
    project, other = demo_plan(), demo_plan()
    store.save(project, "alice", 0)
    store.save(other, "alice", 0)
    project.profile.name = "Synthetic revised project"
    store.save(project, "alice", 1)
    with pytest.raises(ValueError, match="saved project"):
        store.delete(project.id, "alice", 0)
    with pytest.raises(ConflictError):
        store.delete(project.id, "bob", 2)
    with pytest.raises(ConflictError):
        store.delete(project.id, "alice", 1)
    assert store.load(project.id, "alice")[1] == 2
    store.delete(project.id, "alice", 2)
    assert [r["id"] for r in store.list("alice")] == [other.id]
    with pytest.raises(ValueError, match="not found"):
        store.load(project.id, "alice")
    with store.engine.connect() as conn:
        assert not conn.execute(select(store.snapshots).where(store.snapshots.c.id == project.id)).all()
        assert len(conn.execute(select(store.snapshots).where(store.snapshots.c.id == other.id)).all()) == 1


def test_failed_history_deletion_rolls_back_project_removal(tmp_path):
    from sqlalchemy import event, select

    store = Store(f"sqlite:///{tmp_path / 'synthetic-delete-rollback.db'}")
    project = demo_plan()
    store.save(project, "alice", 0)

    def reject_history_delete(conn, cursor, statement, parameters, context, executemany):
        if statement.startswith("DELETE FROM snapshots"):
            raise RuntimeError("Synthetic history-deletion failure")

    event.listen(store.engine, "before_cursor_execute", reject_history_delete)
    try:
        with pytest.raises(RuntimeError, match="Synthetic"):
            store.delete(project.id, "alice", 1)
    finally:
        event.remove(store.engine, "before_cursor_execute", reject_history_delete)
    assert store.load(project.id, "alice")[1] == 1
    with store.engine.connect() as conn:
        assert len(conn.execute(select(store.snapshots).where(store.snapshots.c.id == project.id)).all()) == 1
