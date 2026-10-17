"""Checkpoint regressions: zero is exclusion, current deployment is optional context."""

from io import BytesIO

import openpyxl
import pytest

from services.calculations import budget, preliminary
from services.demo import demo_plan
from services.evidence import load_records, map_role
from services.export import export_excel
from services.hybrid import COMMON_ROLES, comparison_rows, save_input_table, validate_comparison
from services.journey import working_draft
from services.models import Plan, Role
from services.storage import Store
from services.workflow import apply_change, missing_for_gate


def entered():
    p = Plan(contract_kind="existing contract", current_deployment_status="known")
    rows = comparison_rows(p)
    for row in rows:
        row.update(
            {"Current quantity / FTE": 1.0, "Proposed quantity / FTE": 0.0, "Quantity basis": "headcount"}
        )
    rows[0].update(
        {
            "Proposed quantity / FTE": 0.5,
            "Quantity basis": "FTE",
            "Duties / deployment notes": "Shared FM also performs FSM duties",
        }
    )
    rows[2]["Proposed quantity / FTE"] = 4.0
    save_input_table(p, "Synthetic comparison", 2028, 2, 2, rows, "Specialist works by separate contractors")
    return p


def test_common_rows_are_blank_not_compulsory_or_autoallocated():
    rows = comparison_rows(Plan())
    assert [r["Role title"] for r in rows] == COMMON_ROLES
    assert all(r["Proposed quantity / FTE"] is None for r in rows)
    with pytest.raises(ValueError, match="unanswered|Enter proposed"):
        validate_comparison(rows)


def test_zero_excludes_fractional_fte_and_current_difference_survive_storage(tmp_path):
    p = entered()
    d = working_draft(p)
    assert {r.label: r.quantity for r in d.roles if r.included} == {"Facility Manager": 0.5, "Technician": 4}
    fsm = next(r for r in d.roles if r.label == "Fire Safety Manager")
    assert fsm.quantity == 0 and not fsm.included and fsm.exclusion_reason
    rows = comparison_rows(p)
    assert rows[0]["Change"] == -0.5 and rows[2]["Change"] == 3
    assert rows[4]["Change"] == -1
    validate_comparison(rows)
    store = Store(f"sqlite:///{tmp_path / 'comparison.db'}")
    store.save(p, "synthetic", 0)
    restored, revision = store.load(p.id, "synthetic")
    assert revision == 1 and comparison_rows(restored) == rows
    assert not restored.confirmations


@pytest.mark.parametrize("status", ["unavailable", "no existing deployment"])
def test_unknown_current_does_not_block_manpower_input(status):
    p = entered()
    p.current_deployment_status = status
    rows = comparison_rows(p)
    assert all(r["Current quantity / FTE"] is None and r["Change"] is None for r in rows)
    validate_comparison(rows)


def test_included_zero_and_unconverted_posts_cannot_be_priced():
    p = demo_plan()
    p.roles[0].quantity = 0
    assert any("quantity" in e for e in missing_for_gate(p, "quantities"))
    with pytest.raises(ValueError, match="positive quantity"):
        budget(p, load_records())
    with pytest.raises(ValueError):
        preliminary(p)
    p.roles[0].quantity = 1
    p.roles[0].quantity_unit = "posts"
    with pytest.raises(ValueError, match="posts"):
        budget(p, load_records())
    p.roles[0].quantity_unit, p.roles[0].quantity = "headcount", 0.5
    with pytest.raises(ValueError, match="FTE"):
        budget(p, load_records())


def test_seniority_custom_combined_titles_not_mapped_to_basic_fm():
    assert map_role("Facility Manager") == "Facilities Manager"
    assert map_role("Senior Property Executive cum FSM") is None
    assert map_role("M&E Engineer") is None
    assert Role(id="custom", label="FM cum FSM", quantity=0, included=False).label == "FM cum FSM"


def test_pricing_revision_preserves_current_and_compiled_manpower_confirmation():
    p = demo_plan()
    p.roles[0].current_quantity = 2
    p.confirmations = {
        "profile": 1,
        "responsibilities": 1,
        "quantities": 1,
        "plan": 1,
        "rates": 1,
        "base": 1,
        "escalation": 1,
    }
    r = p.roles[0]
    apply_change(
        p, "role", {**r.model_dump(mode="json"), "rate": 9000}, "Synthetic rate edit", "Synthetic officer"
    )
    assert "plan" in p.confirmations and "rates" not in p.confirmations
    assert next(r for r in p.roles if r.id == "fm").current_quantity == 2


def test_excel_keeps_comparison_basis_and_zero_exclusion():
    p = demo_plan()
    p.current_deployment_status = "known"
    p.roles[0].current_quantity, p.roles[0].quantity_unit = 2, "headcount"
    wb = openpyxl.load_workbook(BytesIO(export_excel(p, load_records())), data_only=True)
    ws = wb["Manpower plan"]
    headers = [c.value for c in ws[4]]
    row = [c.value for c in ws[5]]
    assert row[headers.index("Current contract quantity / FTE")] == 2
    assert row[headers.index("Change from current (same basis)")] == -1
    assert row[headers.index("Quantity unit")] == "headcount"
    assert len(wb.sheetnames) == 9 and sum(len(ws._charts) for ws in wb) == 2
