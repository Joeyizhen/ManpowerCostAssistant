"""Synthetic decision journeys and independent salary-component checks; no paid calls."""

from datetime import date
from io import BytesIO

import openpyxl
import pytest

from services.calculations import budget
from services.decision_basis import basis_errors, meaningful, record_basis, retained_basis
from services.demo import demo_plan
from services.evidence import load_records
from services.models import Plan, PolicyAdjustment, Profile, PWMConfig, Role
from services.pwm import catalog, get_schedule
from services.storage import Store
from services.workflow import apply_change, findings


def pwm_plan():
    schedule = next(
        s
        for s in catalog()["schedules"]
        if s["sector"] == "cleaning"
        and "Group 1" in s["category"]
        and s["job_level"] == "General / indoor cleaner"
    )
    config = PWMConfig(
        schedule_digest=get_schedule(schedule["id"])["digest"],
        schedule_id=schedule["id"],
        covered_quantity=2,
        starting_salary=2000,
        salary_source="Synthetic salary-component quotation Q1",
        applicability_confirmed=True,
        applicability_basis="Two included local full-time indoor cleaners confirmed against the selected job category",
        remainder_flat_confirmed=True,
        oncost_rate=0.1,
        oncost_source="Synthetic evidenced employer cost rate, not a legal default",
        rationale="Apply published salary movements to the two cleaners while holding the remaining quoted charge unchanged",
    )
    return Plan(
        contract_kind="existing contract",
        current_deployment_status="known",
        profile=Profile(
            description="Synthetic office with two cleaning zones",
            start=date(2026, 1, 1),
            base_years=1,
            option_years=1,
        ),
        roles=[
            Role(
                id="clean",
                label="Cleaning team (manpower only)",
                quantity=3,
                current_quantity=3,
                quantity_unit="headcount",
                rate=3000,
                escalation=[0, 0],
                pwm=config,
            ),
            Role(
                id="fm",
                label="Facility Manager",
                quantity=1,
                current_quantity=1,
                quantity_unit="headcount",
                rate=8000,
                escalation=[0, 0.02],
            ),
        ],
    )


def test_calendar_month_salary_ratio_and_mixed_roles_reconcile_independently():
    p = pwm_plan()
    result = budget(p, [])
    rows = [r for r in result["rows"] if r["role_id"] == "clean"]
    # 2000 × (2080 / 1910) = 2178.01; six months × 2 workers × 178.01.
    assert rows[0]["pwm_salary_cost"] == 2136.12
    assert rows[0]["pwm_oncost_cost"] == 213.61
    assert rows[0]["annual_cost"] == 108000 + 2136.12 + 213.61
    assert rows[1]["pwm_salary_cost"] == 6408.36
    assert len(rows[0]["pwm_months"]) == 12
    assert rows[0]["pwm_months"][0]["adjusted_salary"] == 2000
    assert rows[0]["pwm_months"][6]["adjusted_salary"] == 2178.01
    assert all(r["pwm_cost"] == 0 for r in result["rows"] if r["role_id"] == "fm")
    assert round(result["base_total"] + result["option_total"], 2) == result["contract_total"]
    assert (
        round(result["waterfall"]["baseline"] + sum(result["waterfall"]["additions"]), 2)
        == result["contract_total"]
    )


@pytest.mark.parametrize(
    "conflict",
    ["percentage", "quote", "policy", "eligibility", "covered", "salary", "partial", "unknown_date"],
)
def test_pwm_rejects_double_counting_and_unconfirmed_component_rules(conflict):
    p = pwm_plan()
    r = p.roles[0]
    if conflict == "percentage":
        r.escalation = [0, 0.02]
    if conflict == "quote":
        r.escalation_mode = "quoted schedule"
    if conflict == "policy":
        r.policies = [
            PolicyAdjustment(
                name="Synthetic duplicate",
                effective_date=date(2026, 7, 1),
                annual_monthly_increment=100,
                affected_component="salary",
                source="Synthetic",
                basis_confirmed=True,
                overlap_reviewed=True,
            )
        ]
    if conflict == "eligibility":
        r.pwm.applicability_confirmed = False
    if conflict == "covered":
        r.pwm.covered_quantity = 4
    if conflict == "salary":
        r.pwm.starting_salary = 3100
    if conflict == "partial":
        p.profile.start = date(2026, 1, 15)
    if conflict == "unknown_date":
        p.profile.start, p.profile.start_year = None, 2026
    with pytest.raises(ValueError):
        budget(p, [])


def test_unpublished_period_is_an_explicit_assumption_not_extended_policy():
    p = pwm_plan()
    p.profile.start = date(2029, 1, 1)
    with pytest.raises(ValueError, match="published PWM period"):
        budget(p, [])
    p.roles[0].pwm.beyond_published = "hold last salary"
    p.roles[
        0
    ].pwm.beyond_reason = "Synthetic officer holds the last published salary pending a new official schedule"
    result = budget(p, [])
    months = result["rows"][0]["pwm_months"]
    assert "Published" in months[0]["basis"]
    assert "Officer assumption" in months[6]["basis"]
    assert months[6]["adjusted_salary"] == 2000


def test_vague_basis_is_rejected_and_one_explanation_can_cover_related_changes():
    p = demo_plan()
    p.contract_kind = "existing contract"
    for r in p.roles:
        r.current_quantity = r.quantity - 1
        r.quantity_reason = "Officer-provided quantity in manpower table"
        r.duties_notes = "Need more manpower"
    assert not meaningful("Need more manpower")
    assert len(basis_errors(p, "quantity", load_records())) == len(p.roles)
    with pytest.raises(ValueError):
        record_basis(
            p,
            "quantity",
            [r.id for r in p.roles],
            "Additional locations",
            "Need more manpower",
            "",
            "officer assumption",
            "Synthetic officer",
        )
    reason = "Three additional buildings require one dedicated officer each for the extended service scope"
    record_basis(
        p,
        "quantity",
        [r.id for r in p.roles],
        "Additional locations",
        reason,
        "Synthetic operating brief",
        "officer assumption",
        "Synthetic officer",
    )
    assert not basis_errors(p, "quantity", load_records())
    second = p.roles[1]
    payload = p.roles[0].model_dump(mode="json")
    payload["quantity"] += 1
    apply_change(p, "role", payload, "Synthetic quantity revision", "Synthetic officer")
    assert retained_basis(p, second, "quantity") is not None
    assert retained_basis(p, next(r for r in p.roles if r.id == payload["id"]), "quantity") is None
    errors = basis_errors(p, "quantity", load_records())
    assert len(errors) == 1 and errors[0].startswith(p.roles[-1].label + ":")


def test_incomplete_basis_draft_saves_without_approving_or_claiming_verified_evidence(tmp_path):
    p = demo_plan()
    p.basis_drafts["quantity"] = {
        "_explanation": "Need more manpower",
        "_classification": "officer assumption",
    }
    store = Store(f"sqlite:///{tmp_path / 'synthetic.db'}")
    store.save(p, "synthetic", 0)
    loaded, _ = store.load(p.id, "synthetic")
    assert loaded.basis_drafts == p.basis_drafts and not loaded.confirmations
    with pytest.raises(ValueError, match="source"):
        record_basis(
            p,
            "quantity",
            [p.roles[0].id],
            "Other",
            "Two additional buildings require two officers for the longer operating coverage",
            "",
            "sourced evidence",
            "Synthetic officer",
        )


def test_basis_gaps_are_distinct_from_invalid_calculation_findings():
    p = demo_plan()
    # An added role needs an operational basis; unchanged existing staffing does not.
    p.roles[0].current_quantity = 0
    p.roles[0].quantity_reason = ""
    p.roles[0].duties_notes = "Need more manpower"
    fs = findings(p, load_records())
    assert any(f.code.startswith("basis_missing_") and not f.blocking for f in fs)
    assert budget(p, load_records())["contract_total"] > 0


def test_pwm_workbook_cached_totals_and_monthly_source_trace():
    from services.export import export_excel

    p = pwm_plan()
    content = export_excel(p, [])
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    assert "PWM monthly trace" in wb.sheetnames
    annual = list(wb["Annual calculations"].iter_rows(min_row=5, max_row=8, values_only=True))
    assert round(sum(row[14] for row in annual), 2) == budget(p, [])["contract_total"]
    assert annual[0][18] == 2349.73
    formula = openpyxl.load_workbook(BytesIO(content), data_only=False)
    assert "+S5" in formula["Annual calculations"]["O5"].value
    assert any("mom.gov.sg" in str(cell.value) for row in wb["PWM monthly trace"] for cell in row)


def test_unrelated_price_changes_preserve_quantity_judgement_and_relevant_changes_reopen():
    from services.workflow import record_decision

    p = demo_plan()
    r = p.roles[0]
    r.current_quantity = 1
    r.quantity = 2
    r.duties_notes = "Additional 4 pax for synthetic support"
    p.contract_kind = "existing contract"
    finding = next(f for f in findings(p, load_records()) if f.code == "quantity_contradiction")
    record_decision(
        p,
        finding,
        "accepted exception",
        "Synthetic officer retains two dedicated workers because the quoted shared support supplies the balance",
        "Synthetic officer",
    )
    payload = next(x for x in p.roles if x.id == r.id).model_dump(mode="json")
    payload["rate"] += 100
    apply_change(p, "role", payload, "Synthetic pricing-only change", "Synthetic officer")
    assert next(f for f in findings(p, load_records()) if f.id == finding.id).officer == "Synthetic officer"
    payload["quantity"] = 3
    apply_change(p, "role", payload, "Synthetic staffing revision", "Synthetic officer")
    assert not next(f for f in findings(p, load_records()) if f.id == finding.id).officer


def test_shared_explicit_building_arithmetic_is_advisory_even_with_recorded_basis():
    from services.manpower_checks import shared_basis_mismatches

    p = demo_plan()
    p.contract_kind = "existing contract"
    for role in p.roles:
        role.current_quantity = role.quantity
    record_basis(
        p,
        "quantity",
        [p.roles[0].id],
        "Additional locations",
        "Two additional buildings require two officers each for dedicated local supervision",
        "Synthetic brief",
        "officer assumption",
        "Synthetic officer",
    )
    checks = shared_basis_mismatches(p)
    assert len(checks) == 1 and "+4" in checks[0]["Check"]
    assert any(f.code.startswith("shared_basis_") and not f.blocking for f in findings(p, load_records()))


def test_contract_start_salary_is_not_uplifted_for_precontract_unpublished_assumptions():
    p = pwm_plan()
    p.profile.start = date(2030, 1, 1)
    p.roles[0].pwm.beyond_published = "officer annual assumption"
    p.roles[0].pwm.beyond_rate = 0.03
    p.roles[
        0
    ].pwm.beyond_reason = (
        "Synthetic employer forecast assumes yearly renewal salary increases pending a new official schedule"
    )
    first = budget(p, [])["rows"][0]["pwm_months"]
    assert first[0]["adjusted_salary"] == 2000
    assert first[6]["adjusted_salary"] == 2060


def test_unapplied_control_snapshot_restores_dates_without_applying_or_claiming_baseline_drafts(monkeypatch):
    import ui_drafts

    p = pwm_plan()
    prefix = f"esc_{p.id}_clean"
    state = {prefix + "_yearly_text": "0, 2", prefix + "_policy_effective": date(2026, 7, 1)}
    monkeypatch.setattr(ui_drafts.st, "session_state", state)
    errors = ["Cleaning team: apply the edited yearly profile and reason before confirming."]
    assert ui_drafts.unsaved(p, errors)
    ui_drafts.capture(p, errors)
    assert not ui_drafts.unsaved(p, errors)
    original_profile = list(p.roles[0].escalation)
    state.clear()
    ui_drafts.restore(p)
    assert state[prefix + "_policy_effective"] == date(2026, 7, 1)
    assert p.roles[0].escalation == original_profile and not p.confirmations
    ui_drafts.capture(p, ["Record or revert the quantity basis draft before confirmation."])
    assert "_ui_controls" not in p.basis_drafts
    assert not ui_drafts.unsaved(p, ["Record or revert the quantity basis draft before confirmation."])


def test_development_import_recovery_keeps_current_models_and_does_not_touch_sessions():
    from types import SimpleNamespace

    from services.runtime_compat import recover_stale_imports

    def old_capture(plan):
        return plan

    current = SimpleNamespace(Plan=SimpleNamespace(model_fields={"decision_bases": object()}))
    modules = {
        "services.models": current,
        "ui_drafts": SimpleNamespace(capture=old_capture),
        "hybrid_inputs": object(),
        "unrelated": object(),
    }
    assert recover_stale_imports(modules)
    assert modules["services.models"] is current and "ui_drafts" not in modules and "unrelated" in modules
    assert not recover_stale_imports(modules)


def test_saved_future_method_remains_pending_when_its_widget_is_not_on_screen(monkeypatch):
    import journey_inputs

    p = pwm_plan()
    p.basis_drafts["future_mode_fm"] = "PWM salary progression"
    monkeypatch.setattr(journey_inputs.st, "session_state", {})
    errors = journey_inputs.pending_control_errors(p)
    assert any(
        error.startswith("Facility Manager: apply the selected future price method") for error in errors
    )
    assert p.roles[1].pwm is None and not p.confirmations
    assert not journey_inputs.pending_control_errors(p, scope="operational")
    assert not journey_inputs.pending_control_errors(p, scope="rates")
    assert p.basis_drafts["future_mode_fm"] == "PWM salary progression"
