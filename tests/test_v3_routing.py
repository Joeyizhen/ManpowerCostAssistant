"""Workbook routing and mocked assistant contracts, without paid model calls."""

import json
from types import SimpleNamespace as NS

import pytest
from test_agent import MockResponses, call, response
from test_budget_basis_pwm import pwm_plan

from services.agent import converse
from services.calculations import budget
from services.cost_routing import pin_route, recorded_route, resolve_route, route_errors, schedule_choices
from services.evidence import load_records
from services.future_groups import general_proposal
from services.journey import working_draft
from services.models import Role
from services.reference_data import root
from services.reference_import import activate, validate_activation
from services.role_catalogue import available_roles, catalogue
from services.workflow import apply_change

pytest_plugins = ["test_prepared_references"]


@pytest.fixture
def routing_source(prepared):
    source = json.loads(json.dumps(prepared))
    source["dataset_id"] = "routing-fixture-v3"
    source["source_name"] = "invented-v3.xlsx"
    source["standard_roles"] = [
        {"Role Title": title, "PWM applies?": flag, "Typical manpower?": typical}
        for title, flag, typical in [
            ("Facility Manager", "No", "Yes"),
            ("Cleaner", "Yes", "Yes"),
            ("Cleaning Supervisor", "Yes", "Yes"),
            ("Security Officer", "Yes", "Yes"),
            ("Senior Security Officer", "Yes", "Yes"),
            ("Security Supervisor", "Yes", "Yes"),
            ("Lift Technician", "Yes", "No"),
            ("Landscape Technician", "No", "No"),
        ]
    ]
    (root() / "snapshots" / (source["dataset_id"] + ".json")).write_text(json.dumps(source))
    activate(source["dataset_id"])
    return source


def test_every_flag_controls_route_without_keywords(routing_source):
    for item in routing_source["standard_roles"]:
        role = Role(id=item["Role Title"], label=item["Role Title"])
        assert resolve_route(role)["route"] == ("pwm" if item["PWM applies?"] == "Yes" else "general")
    assert resolve_route(Role(id="u", label="Unlisted security supervisor"))["route"] == "unresolved"
    assert resolve_route(Role(id="u", label="Cleaner cum Facilities Manager"))["route"] == "unresolved"
    assert resolve_route(Role(id="l", label="Landscape Technician"))["route"] == "general"
    assert not schedule_choices(Role(id="l", label="Landscape Technician"))


def test_catalogue_and_aliases_do_not_duplicate_positions(routing_source):
    result = catalogue()
    assert "Lift Technician" in result["additional"] and "Cleaner" in result["typical"]
    assert not set(result["typical"]) & set(result["additional"])
    from services.models import Plan

    plan = Plan(roles=[Role(id="fm", label="FM", quantity=1, quantity_unit="headcount")])
    assert "Facility Manager" not in available_roles(plan)["typical"]


def test_policy_levels_and_unknown_flags_are_explicit(routing_source):
    for title, level in [
        ("Security Officer", "Security officer"),
        ("Senior Security Officer", "Senior security officer"),
        ("Security Supervisor", "Security supervisor"),
        ("Cleaning Supervisor", "Supervisor"),
    ]:
        schedules = schedule_choices(Role(id=title, label=title))
        assert schedules and all(s["job_level"] == level for s in schedules)
    missing = Role(id="u", label="Not in taxonomy")
    assert route_errors(missing, True)
    with pytest.raises(ValueError, match="no default"):
        pin_route(missing)


def test_mixed_budget_no_overlap_and_calendar_months(routing_source):
    plan = pwm_plan()
    plan.roles[0].label = "Cleaner"
    pin_route(plan.roles[0])
    general = Role(
        id="fm",
        label="Facility Manager",
        quantity=1,
        quantity_unit="headcount",
        rate=4000,
        escalation=[0, 0.03],
    )
    pin_route(general)
    plan.roles = [plan.roles[0], general]
    result = budget(plan, load_records())
    cleaner = [r for r in result["rows"] if r["role_id"] == "clean"]
    assert all(r["annual_before_policy"] == r["flat_baseline"] for r in cleaner)
    assert all(len(r["pwm_months"]) == 12 for r in cleaner)
    assert sum(r["pwm_cost"] for r in result["rows"]) > 0
    assert sum(r["excess_over_baseline"] for r in result["rows"] if r["role_id"] == "fm") == pytest.approx(
        1440
    )
    bad = plan.model_copy(deep=True)
    bad.roles[0].escalation = [0, 0.03]
    with pytest.raises(ValueError, match="stacked"):
        budget(bad, load_records())
    with pytest.raises(ValueError, match="PWM workers"):
        general_proposal(plan, load_records(), ["clean"], 3, "", "", False, "Officer")


def test_snapshot_pins_survive_activation_and_old_budgets_are_not_rewritten(routing_source, prepared):
    plan = pwm_plan()
    before = plan.model_dump(mode="json")
    original_total = budget(plan, load_records())["contract_total"]
    assert plan.model_dump(mode="json") == before
    assert budget(plan, load_records())["contract_total"] == original_total
    role = Role(id="fm", label="Facility Manager")
    pin_route(role)
    activate(prepared["dataset_id"])
    assert recorded_route(role)["dataset_id"] == routing_source["dataset_id"]
    assert recorded_route(role)["route"] == "general"


def test_invalid_snapshot_cannot_activate(prepared):
    invalid = json.loads(json.dumps(prepared))
    invalid["readiness_report"]["structural_errors"] = ["Missing PWM flag"]
    with pytest.raises(ValueError, match="activation refused"):
        validate_activation(invalid)


def test_new_role_routes_are_enforced_by_calculator(routing_source):
    plan = pwm_plan()
    item = plan.roles[0].model_copy(deep=True)
    item.id, item.label = "new-fm", "Facility Manager"
    apply_change(plan, "role", item.model_dump(mode="json"), "Officer added a proposal role", "Officer")
    added = next(r for r in plan.roles if r.id == "new-fm")
    assert added.future_routing["pwm_flag"] == "No"
    with pytest.raises(ValueError, match="PWM applies.*No"):
        budget(plan, load_records())


def test_blank_flag_does_not_default_to_no(routing_source, monkeypatch):
    routing_source["standard_roles"][0]["PWM applies?"] = None
    monkeypatch.setattr("services.cost_routing.active_snapshot", lambda: routing_source)
    assert resolve_route(Role(id="fm", label="Facility Manager"))["route"] == "unresolved"


def test_required_flags_and_blank_rows_in_import(tmp_path):
    from openpyxl import load_workbook
    from test_prepared_references import workbook

    from services.reference_import import build_snapshot

    path = workbook(tmp_path / "invalid.xlsx")
    book = load_workbook(path)
    book["Standard Roles"].append([None] * 4)
    book.save(path)
    clean = build_snapshot(path)
    assert len(clean["standard_roles"]) == 1 and not clean["readiness_report"]["structural_errors"]
    book["Standard Roles"].cell(2, 4, "Unknown")
    book.save(path)
    invalid = build_snapshot(path)
    assert invalid["readiness_report"]["structural_errors"]
    with pytest.raises(ValueError, match="activation refused"):
        validate_activation(invalid)


def test_prepared_pwm_routing_ui(routing_source, monkeypatch, tmp_path):
    from test_ui import blank

    at = blank(monkeypatch, tmp_path)
    plan = pwm_plan()
    plan.roles[0].label = "Cleaner"
    plan.ui_step = 4
    plan.confirmations = {
        gate: plan.version for gate in ("profile", "responsibilities", "quantities", "plan", "rates", "base")
    }
    at.session_state["plan"] = plan
    at.run()
    assert not at.exception
    worker_control = next(m for m in at.multiselect if m.label == "Workers to check for PWM coverage")
    assert worker_control.disabled
    assert at.session_state["plan"].roles[0].pwm is not None
    assert at.session_state["plan"].roles[1].rate == plan.roles[1].rate
    assert any("PWM salary progression" in m.value for m in at.markdown)
    # Only the general FM role appears in the general group, not the cleaner.
    general = next(d for d in at.dataframe if "Manpower" in d.value.columns)
    assert set(general.value["Manpower"]) == {"Facility Manager"}


def test_assistant_mismatch_interaction_is_proposal_only(routing_source):
    plan = pwm_plan()
    plan.roles[0].label = "Facility Manager"  # Existing proposal attempts PWM on a No role.
    plan.confirmations = {"plan": plan.version, "rates": plan.version}
    before = plan.model_dump(mode="json")
    sdk = MockResponses(
        [
            response(
                [
                    call(
                        "check_future_cost_method",
                        {"role_title": "Facility Manager", "requested_method": "PWM salary progression"},
                    )
                ]
            ),
            response(
                [
                    call(
                        "propose_escalation",
                        {
                            "role_id": "clean",
                            "yearly_changes": [0, 0.1],
                            "rationale": "Officer may accept the supported general reference; PWM conflicts with workbook No flag.",
                            "mode": "assumption",
                            "schedule_id": None,
                        },
                        "proposal",
                    )
                ]
            ),
            response(
                text="PWM applies? is No for Facility Manager. The general reference is 10%. Review the proposed replacement before accepting."
            ),
        ]
    )
    converse(
        plan,
        load_records(),
        "My proposed Facility Manager allowance uses PWM. Can you review it?",
        NS(responses=sdk),
        "mock-only",
        "Escalation",
        review_existing_proposal=True,
    )
    checked = plan.tool_audit[0]["result"]
    assert checked["conflict"] and not checked["plan_changed"]
    assert checked["general_reference"]["average"] == pytest.approx(0.1)
    assert plan.roles[0].model_dump(mode="json") == before["roles"][0]
    proposed = next(r for r in working_draft(plan).roles if r.id == "clean")
    assert proposed.pwm is None and proposed.escalation == [0, 0.1]
    assert proposed.future_routing["pwm_flag"] == "No"
    apply_change(
        plan, "role", proposed.model_dump(mode="json"), "Officer accepts the reviewed correction", "Officer"
    )
    assert plan.roles[0].pwm is None and plan.confirmations.get("rates") is not None
    assert (
        budget(plan, load_records())["contract_total"] != budget(pwm_plan(), load_records())["contract_total"]
    )
    assert before["roles"][0]["pwm"] is not None  # Original proposal copy remains intact.
