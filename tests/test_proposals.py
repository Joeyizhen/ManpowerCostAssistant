"""Independent Track 2 capture, evidence preservation, calculations and approval."""

import io
import zipfile

import openpyxl
import pytest
import xlsxwriter
from sqlalchemy import create_engine

from proposal_inputs import current_rows, unapplied_figures
from services.evidence import load_records, rate_benchmark
from services.models import Plan
from services.proposal_export import management_report, workbook
from services.proposals import (
    Proposal,
    ProposalStore,
    apply_yearly,
    assessment,
    capture_original,
    comparison,
    confirm_review,
    extract_document,
    interpret_text,
    normalise_table,
    original_signature,
    restore_copy,
    update_reviewed,
)
from services.storage import ConflictError
from services.workflow import confirm, record_decision


def sample(rate=8000, annual=0, quantity=1):
    p = Proposal()
    rows = normalise_table(
        [
            {
                "Role": "Facility Manager",
                "Quantity": quantity,
                "Unit": "headcount",
                "Monthly rate": rate,
                "Annual increase (%)": annual,
                "Remarks": "Weekday office coordination workload",
            }
        ],
        "Synthetic quotation Q1",
    )
    capture_original(
        p,
        {
            "name": "Synthetic proposal review",
            "description": "Weekday office coordination; technical works separately procured",
            "start": "2028-01-01",
            "start_year": 2028,
            "base_years": 1,
            "option_years": 1,
        },
        rows,
        True,
    )
    return p


def test_original_is_immutable_across_accepted_revisions_and_owner_saves(tmp_path):
    p = sample()
    before = p.original.model_dump_json()
    store = ProposalStore(create_engine(f"sqlite:///{tmp_path / 'owned.db'}"))
    assert store.save(p, "owner", 0) == 1
    row = {**p.original_rows[0], "Quantity": 2}
    update_reviewed(p, p.plan.profile.model_dump(mode="json"), [row], "Officer")
    assert p.original.model_dump_json() == before
    assert original_signature(p) == p.original_digest
    assert comparison(p, load_records())["difference"] == 192000
    assert store.save(p, "owner", 1) == 2
    with pytest.raises(ValueError, match="access denied"):
        store.load(p.id, "other")
    with pytest.raises(ConflictError):
        store.save(p, "owner", 1)
    p.original.roles[0].quantity = 99
    with pytest.raises(ValueError, match="cannot be overwritten"):
        store.save(p, "owner", 2)


def test_capture_preserves_missing_units_rates_and_explicit_zero():
    p = sample()
    p.plan.roles[0].quantity_unit = "unspecified"
    state = assessment(p, load_records())
    assert state["result"] is None
    assert any(f.code == "proposal_unit" for f in state["blockers"])
    p.plan.roles[0].quantity_unit = "headcount"
    p.plan.roles[0].rate = None
    assert assessment(p, load_records())["result"] is None
    p = sample(annual=0)
    assert comparison(p, load_records())["reviewed"]["annual_totals"] == [96000, 96000]
    p.plan.roles[0].escalation = []
    assert assessment(p, load_records())["result"] is None


def test_conservative_free_text_and_csv_keep_titles_sources_and_unknowns():
    rows, notes = interpret_text(
        "one senior FM, four M&E engineers and two technicians. Use 3% escalation for six years"
    )
    assert [r["Quantity"] for r in rows] == [1, 4, 2]
    assert rows[0]["Role"] == "senior FM"
    assert all(r["Unit"] == "unspecified" and r["Monthly rate"] is None for r in rows)
    assert all(r["Annual increase (%)"] == 3 for r in rows)
    assert notes and all("characters" in r["Source location"] for r in rows)
    doc = extract_document(
        "proposal.csv",
        b"Role,Quantity,Unit,Monthly rate,Annual increase (%)\nFM / Events Manager,0.5,FTE,6000,0\n",
    )
    assert doc["rows"][0]["Role"] == "FM / Events Manager"
    assert doc["rows"][0]["Unit"] == "FTE" and doc["bytes_base64"]


def test_xlsx_docx_pdf_extraction_keeps_locations_without_execution():
    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out, {"in_memory": True})
    sheet = wb.add_worksheet("Manpower")
    sheet.write_row(0, 0, ["Role", "Quantity", "Unit", "Monthly rate"])
    sheet.write_row(1, 0, ["Facility Manager", 1, "headcount", 8000])
    wb.close()
    doc = extract_document("schedule.xlsx", out.getvalue())
    assert doc["rows"][0]["Monthly rate"] == 8000
    assert "sheet1.xml" in doc["rows"][0]["Source location"]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>one senior FM</w:t></w:r></w:p></w:body></w:document>',
        )
    assert extract_document("idea.docx", out.getvalue())["rows"][0]["Role"] == "senior FM"
    from pypdf import PdfWriter

    out = io.BytesIO()
    pdf = PdfWriter()
    pdf.add_blank_page(width=300, height=300)
    pdf.write(out)
    doc = extract_document("scanned.pdf", out.getvalue())
    assert not doc["rows"] and "manual transcription" in doc["extraction_note"]


def test_supplied_arithmetic_is_preserved_and_requires_explicit_decision():
    p = sample()
    p.original_claimed_total = 1
    p.original_digest = original_signature(p)
    state = assessment(p, load_records())
    flag = next(f for f in state["unresolved"] if f.code == "proposal_math")
    with pytest.raises(ValueError, match="material findings"):
        confirm_review(p, load_records(), "Officer", "Retain original office deployment")
    record_decision(
        p.plan,
        flag,
        "accepted exception",
        "Supplied total omitted the second contract year; accept the recalculated two-year EPV",
        "Officer",
    )
    assert not assessment(p, load_records())["unresolved"]
    assert p.original_claimed_total == 1


def test_review_confirmation_is_distinct_from_final_and_invalidates_after_change():
    p = sample()
    confirm_review(
        p,
        load_records(),
        "Officer",
        "Weekday office coordination workload retained; technical duties separately contracted",
    )
    assert "validation" in p.plan.confirmations and "final" not in p.plan.confirmations
    confirm(p.plan, "final", "Officer", load_records())
    assert p.plan.final_officer == "Officer"
    row = {**p.original_rows[0], "Quantity": 2}
    update_reviewed(p, p.plan.profile.model_dump(mode="json"), [row], "Officer")
    assert "final" not in p.plan.confirmations and "rates" not in p.plan.confirmations
    with pytest.raises(ValueError):
        confirm(p.plan, "final", "Officer", load_records())


def test_revised_price_does_not_relabel_the_original_quote_as_support():
    p = sample()
    revised = {**p.original_rows[0], "Monthly rate": 9000}
    update_reviewed(p, p.plan.profile.model_dump(mode="json"), [revised], "Officer")
    assert p.plan.roles[0].rate_reason == ""
    assert assessment(p, load_records())["basis_actions"]
    revised["Pricing basis"] = "Supplier revised quotation Q2 for the changed inclusive scope"
    update_reviewed(p, p.plan.profile.model_dump(mode="json"), [revised], "Officer")
    assert not assessment(p, load_records())["basis_actions"]
    assert p.original.roles[0].rate == 8000


def test_revised_units_and_contract_dates_preserve_the_original_pricing_basis():
    p = sample()
    row = {**p.original_rows[0], "Unit": "FTE", "Quantity": 0.5}
    update_reviewed(p, p.plan.profile.model_dump(mode="json"), [row], "Officer")
    assert p.plan.roles[0].rate_reason == ""
    assert assessment(p, load_records())["basis_actions"]
    p = sample()
    profile = p.plan.profile.model_dump(mode="json")
    profile.update(start="2029-01-01", start_year=2029)
    update_reviewed(p, profile, p.original_rows, "Officer")
    assert p.plan.roles[0].benchmark_year == 2028 and p.plan.roles[0].rate == 8000
    assert any("reference-year" in text for text in assessment(p, load_records())["basis_actions"])
    assert p.original.profile.start.year == 2028


def test_private_quotes_do_not_change_reference_cohort_or_stack_escalation():
    p = sample()
    rid = p.plan.roles[0].id
    records = load_records()
    before = rate_benchmark(records, "Facilities Manager", 2028)
    apply_yearly(p, rid, None, [8000, 10000], "Synthetic quotation Q2, pages 2-3", "Officer")
    assert comparison(p, records)["reviewed"]["annual_totals"] == [96000, 120000]
    assert p.plan.roles[0].escalation == [0, 0]
    assert rate_benchmark(records + p.quote_records, "Facilities Manager", 2028) == before
    assert p.original.roles[0].escalation_mode == "assumption"
    p.plan.profile.start = p.plan.profile.start.replace(month=2)
    with pytest.raises(ValueError, match="1 January"):
        apply_yearly(p, rid, None, [8000, 10000], "Supplied quote", "Officer")


def test_custom_years_compound_and_exports_reconcile_without_source_formula_execution():
    p = sample()
    apply_yearly(p, p.plan.roles[0].id, [0, 5], None, "Supplier renewal allowance", "Officer")
    result = comparison(p, load_records())
    assert result["reviewed"]["contract_total"] == 196800
    data = openpyxl.load_workbook(io.BytesIO(workbook(p, load_records())), data_only=True)
    total = next(row for row in data["Budget summary"].values if row[0] == "Annual overall totals")
    assert total[3] == 196800
    assert result["original"]["contract_total"] == 192000
    assert {"Original proposal", "Reviewed proposal", "Findings and decisions", "Source locations"} <= set(
        data.sheetnames
    )
    p.name = "<script>alert(1)</script>"
    report = management_report(p, load_records())
    assert "<script>" not in report and "&lt;script&gt;" in report


def test_track1_store_and_plan_schema_are_not_repurposed(tmp_path):
    from services.storage import Store

    store = Store(f"sqlite:///{tmp_path / 'combined.db'}")
    original = Plan()
    store.save(original, "owner", 0)
    p = sample()
    proposals = ProposalStore(store.engine)
    proposals.save(p, "owner", 0)
    saved, revision = store.load(original.id, "owner")
    assert saved.model_dump_json() == original.model_dump_json() and revision == 1
    assert not store.list("other") and not proposals.list("other")


def test_unapplied_draft_values_survive_storage_without_changing_the_recorded_plan(tmp_path):
    p = sample()
    p.drafts["review_rows"] = current_rows(p)
    p.drafts["review_profile"] = p.plan.profile.model_dump(mode="json")
    assert not unapplied_figures(p)
    p.drafts["review_rows"][0]["Monthly rate"] = 9000
    assert unapplied_figures(p)
    store = ProposalStore(create_engine(f"sqlite:///{tmp_path / 'drafts.db'}"))
    store.save(p, "owner", 0)
    saved, _ = store.load(p.id, "owner")
    assert unapplied_figures(saved)
    assert saved.drafts["review_rows"][0]["Monthly rate"] == 9000
    assert saved.plan.roles[0].rate == 8000 and not saved.plan.confirmations
    assert saved.original_digest == p.original_digest


def test_restore_is_a_separate_unapproved_review_with_the_same_original():
    p = sample()
    confirm_review(
        p,
        load_records(),
        "Officer",
        "Retain weekday coordination workload; specialist duties separately procured",
    )
    confirm(p.plan, "final", "Officer", load_records())
    original = p.original.model_dump_json()
    p.drafts["controls"] = {f"proposal_rate_{p.id}_{p.plan.roles[0].id}_{p.plan.version}_exact": 9000}
    copied = restore_copy(p)
    assert copied.id != p.id and copied.plan.id != p.plan.id
    assert not copied.plan.confirmations and not copied.plan.final_date
    assert copied.original.model_dump_json() == original
    assert original_signature(copied) == copied.original_digest
    assert comparison(copied, load_records())["reviewed"]["contract_total"] == 192000
    assert any(str(copied.id) in key and value == 9000 for key, value in copied.drafts["controls"].items())


def test_single_year_needs_no_extra_future_assumption_and_combined_titles_stay_intact():
    p = Proposal()
    row = normalise_table(
        [{"Role": "Facility Manager", "Quantity": 1, "Unit": "headcount", "Monthly rate": 8000}],
        "Synthetic Q1",
    )
    capture_original(
        p,
        {
            "name": "One year",
            "description": "Weekday office coordination",
            "start_year": 2028,
            "base_years": 1,
            "option_years": 0,
        },
        row,
        True,
    )
    assert assessment(p, load_records())["result"]["contract_total"] == 96000
    assert p.plan.roles[0].escalation == [0]
    parsed, _ = interpret_text("one FM / Events Manager. Use -2% escalation")
    assert parsed[0]["Role"] == "FM / Events Manager"
    assert parsed[0]["Annual increase (%)"] == -2


def test_unsupported_original_quotes_remain_capturable_and_overlap_is_material():
    p = Proposal()
    rows = normalise_table(
        [
            {
                "Role": "M&E Engineer",
                "Quantity": 1,
                "Unit": "headcount",
                "Year 1 rate": 8000,
                "Year 2 rate": 9000,
                "Annual increase (%)": 3,
            }
        ],
        "Synthetic quotation",
    )
    capture_original(
        p,
        {
            "name": "Incomplete quote interpretation",
            "description": "Office engineering support",
            "start_year": 2028,
            "base_years": 1,
            "option_years": 1,
        },
        rows,
        True,
    )
    assert p.original and p.original_rows[0]["Yearly rates"] == [8000, 9000]
    assert p.drafts["quote_interpretation_issues"]
    assert comparison(p, load_records())["original"] is None
    assert any(f.code == "proposal_quote_overlap" for f in assessment(p, load_records())["unresolved"])


def test_track2_offline_capture_check_results_confirm_and_resume(monkeypatch, tmp_path):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'journey.db'}")
    monkeypatch.setenv("DEPLOYMENT_MODE", "local")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("OPENAI_MODEL", "")
    at = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=50).run()
    original_track1 = at.session_state["plan"].model_dump_json()
    at.button(key="activity_budgeting").click().run()
    at.button(key="choose_track_2").click().run()
    assert not at.exception
    p = sample()
    at.session_state["proposal"] = p
    at.run()
    assert not at.exception
    assert at.session_state["proposal"].plan.ui_step == 1
    at.button(key="t2_confirm_deployment_" + p.id).click().run()
    assert not at.exception
    assert at.session_state["proposal"].plan.ui_step == 3
    role = at.session_state["proposal"].plan.roles[0]
    prefix = f"rate_{p.plan.id}_{role.id}_{role.benchmark_year}"
    at.number_input(key=prefix + "_number").set_value(9000.0).run()
    assert not at.exception
    assert at.session_state["proposal"].plan.roles[0].rate == 9000
    assert at.session_state["proposal"].original.roles[0].rate == 8000
    note = next((n for n in at.text_area if n.label.startswith("Short pricing reason")), None)
    if note:
        note.set_value(
            "Revised supplier quotation for weekday coordination with the same inclusive basis"
        ).run()
    at.button(key="confirm_rates").click().run()
    assert not at.exception
    assert at.session_state["proposal"].plan.ui_step == 4
    at.button(key="confirm_escalation").click().run()
    assert not at.exception
    assert [x.label for x in at.tabs] == [
        "Budget",
        "Scenarios",
        "Manpower",
        "Assumptions & sources",
        "Changes from your original proposal",
    ]
    assert "final" not in at.session_state["proposal"].plan.confirmations
    at.button(key="proposal_final_" + p.id).click().run()
    assert not at.exception
    assert "final" in at.session_state["proposal"].plan.confirmations
    assert at.session_state["plan"].model_dump_json() == original_track1
    assert at.session_state["proposal"].original_digest == p.original_digest
