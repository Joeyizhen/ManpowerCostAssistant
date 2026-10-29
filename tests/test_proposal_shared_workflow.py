"""Isolated proposal extraction, rule checks, shared decisions and mocked interaction."""

import hashlib
import io
import json

import openpyxl
import pytest
from test_budget_basis_pwm import pwm_plan

from services.calculations import budget
from services.evidence import load_records
from services.proposal_export import workbook
from services.proposal_intake import attach_assumptions, load_example, template_bytes
from services.proposal_workflow import (
    accept_correction,
    confirm_deployment,
    correction_preview,
    material_checks,
    next_decision,
    resolve_conflict,
)
from services.proposals import (
    Proposal,
    ProposalStore,
    assessment,
    capture_original,
    comparison,
    extract_document,
    normalise_table,
    original_signature,
)
from services.reference_data import root
from services.reference_import import activate
from services.workflow import apply_change, confirm

pytest_plugins = ["test_prepared_references"]


@pytest.fixture
def proposal_source(prepared):
    source = json.loads(json.dumps(prepared))
    source["dataset_id"] = "proposal-shared-fixture-v1"
    source["source_name"] = "invented-proposal-reference.xlsx"
    source["standard_roles"] = [
        {"Role Title": name, "Typical manpower?": "Yes", "PWM applies?": flag}
        for name, flag in [
            ("Facility Manager", "No"),
            ("Facilities Manager", "No"),
            ("Admin Executive", "No"),
            ("Technician", "No"),
            ("Cleaner", "Yes"),
        ]
    ]
    (root() / "snapshots" / (source["dataset_id"] + ".json")).write_text(json.dumps(source))
    activate(source["dataset_id"])
    load_records()
    return source


def captured_example():
    p = load_example()
    capture_original(p, p.input_profile, p.input_rows, True)
    return p


def complete_example(p, records):
    technician = next(r for r in p.plan.roles if r.label == "Technician")
    preview = correction_preview(p, technician.id, records)
    accept_correction(p, preview, records, "Officer")
    confirm_deployment(p, records, "Officer")
    for gate in ("rates", "base"):
        confirm(p.plan, gate, "Officer", records)
    cleaner = next(r for r in p.plan.roles if r.label == "Cleaner")
    item = cleaner.model_copy(deep=True)
    item.pwm = (
        pwm_plan()
        .roles[0]
        .pwm.model_copy(update={"oncost_rate": 0.0, "oncost_source": "", "starting_salary": 2300.0})
    )
    item.escalation = [0, 0]
    item.escalation_reason = (
        "Officer accepted the source salary component and published schedule treatment; remainder held flat."
    )
    apply_change(
        p.plan, "role", item.model_dump(mode="json"), "Officer accepted salary-only PWM setup", "Officer"
    )
    confirm(p.plan, "escalation", "Officer", records)
    return p


def test_blank_and_example_templates_are_readable_and_have_no_preloaded_diagnosis(proposal_source):
    blank = extract_document("blank.xlsx", template_bytes())
    assert blank["rows"] == [] and blank["profile"]["base_years"] is None
    assert blank["profile"]["option_years"] is None
    p = load_example()
    assert p.original is None and not p.plan.roles and not p.plan.confirmations
    assert not p.plan.decisions and not p.plan.pending
    assert {r["Proposed method"] for r in p.input_rows} == {"PWM"}
    assert {r["Role"] for r in p.input_rows} == {"Technician", "Cleaner"}
    assert p.input_profile["start"] == "2027-01-01"
    book = openpyxl.load_workbook(io.BytesIO(template_bytes(True)), data_only=True)
    assert book.sheetnames == ["Project", "Proposed manpower", "Proposed escalation"]
    values = " ".join(
        str(v) for sheet in book for row in sheet.values for v in row if v is not None
    ).casefold()
    assert "incorrect" not in values and "correction" not in values and "pwm applies?" not in values


def test_deployment_only_and_explicit_fte_are_accepted(proposal_source):
    p = Proposal()
    rows = normalise_table(
        [{"Role": "Technician", "Quantity": 0.5, "Unit": "FTE"}], "Invented deployment only"
    )
    capture_original(
        p,
        {
            "name": "Deployment only",
            "description": "Shared weekday equipment coordination across offices",
            "start_year": 2027,
            "base_years": 1,
            "option_years": 1,
        },
        rows,
        True,
    )
    assert p.plan.roles[0].rate is None and p.plan.roles[0].escalation == []
    assert p.plan.roles[0].quantity == 0.5 and p.plan.roles[0].quantity_unit == "FTE"
    assert comparison(p, load_records())["original"] is None
    confirm_deployment(p, load_records(), "Officer")
    assert next_decision(p) == 3
    assert not p.plan.confirmations.get("rates")


def test_wrong_pwm_imports_then_explicit_correction_preserves_original(proposal_source):
    p = captured_example()
    original = p.original.model_dump_json()
    digest = original_signature(p)
    findings = material_checks(p, load_records())
    assert [(f["Role"], f["kind"]) for f in findings] == [("Technician", "wrong_method")]
    assert assessment(p, load_records())["result"] is None
    rid = next(r.id for r in p.plan.roles if r.label == "Technician")
    preview = correction_preview(p, rid, load_records())
    assert preview["Selected annual increase (%)"] == pytest.approx(10)
    assert preview["Proposed role contract cost (SGD)"] == 504000
    assert next(r for r in p.plan.roles if r.id == rid).escalation == []
    accept_correction(p, preview, load_records(), "Officer")
    assert next(r for r in p.plan.roles if r.id == rid).escalation == [0, 0.1]
    assert p.original.model_dump_json() == original and original_signature(p) == digest
    assert p.original_rows[0]["Proposed method"] == "PWM"
    assert not [f for f in material_checks(p, load_records()) if f["kind"] == "wrong_method"]
    assert comparison(p, load_records())["original"] is None
    assert comparison(p, load_records())["difference"] is None


def test_correction_preview_is_stale_after_a_deployment_edit(proposal_source):
    p = captured_example()
    role = next(r for r in p.plan.roles if r.label == "Technician")
    preview = correction_preview(p, role.id, load_records())
    payload = role.model_dump(mode="json")
    payload["quantity"] = 5
    apply_change(p.plan, "role", payload, "Officer revised deployment", "Officer")
    with pytest.raises(ValueError, match="Preview this correction again"):
        accept_correction(p, preview, load_records(), "Officer")
    assert p.plan.roles[0].escalation == []
    assert original_signature(p) == p.original_digest


def test_shared_and_individual_assumptions_conflict_until_accepted(proposal_source):
    rows = normalise_table(
        [
            {
                "Role": "Technician",
                "Quantity": 4,
                "Unit": "headcount",
                "Monthly rate": 5000,
                "Remarks": "Four weekday technical responders",
            }
        ],
        "Invented quote",
    )
    base = {
        "method": "General annual increase",
        "yearly_increases": None,
        "yearly_rates": None,
        "note": "Proposal allowance",
        "source": "Proposal escalation row",
    }
    attach_assumptions(
        rows,
        [
            {**base, "target": "All non-PWM roles", "percent": 3},
            {**base, "target": "Technician", "percent": 4},
        ],
    )
    p = Proposal()
    capture_original(
        p,
        {
            "description": "Weekday office equipment response",
            "start_year": 2027,
            "base_years": 1,
            "option_years": 1,
        },
        rows,
        True,
    )
    assert rows[0]["Assumption conflict"] and p.plan.roles[0].escalation == []
    assert any(f["kind"] == "conflict" for f in material_checks(p, load_records()))
    resolve_conflict(p, p.plan.roles[0].id, rows[0]["Supplied assumptions"][1], load_records(), "Officer")
    assert p.plan.roles[0].escalation == [0, 0.04]
    assert not material_checks(p, load_records())
    assert len(p.original_rows[0]["Supplied assumptions"]) == 2


def test_structural_error_names_cell_but_questionable_method_imports(proposal_source):
    book = openpyxl.load_workbook(io.BytesIO(template_bytes(True)))
    book["Proposed manpower"]["B5"] = "not a number"
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(ValueError, match="Proposed manpower!B5"):
        extract_document("bad.xlsx", stream.getvalue())
    p = captured_example()
    assert p.interpretation_confirmed
    assert any(f.code == "proposal_method" for f in assessment(p, load_records())["blockers"])


def test_mixed_budget_export_and_original_comparison_reconcile(proposal_source):
    p = complete_example(captured_example(), load_records())
    result = budget(p.plan, load_records())
    assert result["contract_total"] == pytest.approx(sum(result["annual_totals"]))
    assert result["base_total"] + result["option_total"] == pytest.approx(result["contract_total"])
    clean = [r for r in result["rows"] if r["role"] == "Cleaner"]
    assert all(r["annual_before_policy"] == r["flat_baseline"] for r in clean)
    assert sum(r["pwm_cost"] for r in clean) > 0
    assert (
        next(r for r in result["rows"] if r["role"] == "Technician" and r["year"] == 2)["monthly_unit_rate"]
        == 5500
    )
    data = openpyxl.load_workbook(io.BytesIO(workbook(p, load_records())), data_only=True)
    assert {
        "Manpower plan",
        "Original proposal",
        "Reviewed proposal",
        "Changes from original",
        "Source locations",
    } <= set(data.sheetnames)
    assert comparison(p, load_records())["original"] is None
    assert original_signature(p) == p.original_digest


def test_published_dates_and_overlap_still_block(proposal_source):
    p = complete_example(captured_example(), load_records())
    cleaner = next(r for r in p.plan.roles if r.label == "Cleaner")
    cleaner.escalation = [0, 0.03]
    with pytest.raises(ValueError, match="stacked"):
        budget(p.plan, load_records())
    cleaner.escalation = [0, 0]
    p.plan.profile.start = p.plan.profile.start.replace(day=15)
    with pytest.raises(ValueError, match="first-of-month"):
        budget(p.plan, load_records())


def test_store_reopen_keeps_original_and_pending_inputs(proposal_source, tmp_path):
    from sqlalchemy import create_engine

    p = captured_example()
    p.drafts["review_rows"] = [{**p.original_rows[0], "Quantity": 7}, p.original_rows[1]]
    store = ProposalStore(create_engine(f"sqlite:///{tmp_path / 'draft.db'}"))
    store.save(p, "owner", 0)
    restored, _ = store.load(p.id, "owner")
    assert restored.drafts["review_rows"][0]["Quantity"] == 7
    assert restored.plan.roles[0].quantity == 4
    assert restored.original_digest == p.original_digest
    assert hashlib.sha256(template_bytes(True)).hexdigest() == restored.original_documents[0]["sha256"]


def test_native_excel_percent_values_keep_their_units(proposal_source):
    book = openpyxl.load_workbook(io.BytesIO(template_bytes(True)))
    sheet = book["Proposed escalation"]
    sheet["B5"] = "General annual increase"
    sheet["C5"] = 0.03
    sheet["C5"].number_format = "0.0%"
    stream = io.BytesIO()
    book.save(stream)
    captured = extract_document("native-percent.xlsx", stream.getvalue())
    assert captured["rows"][0]["Annual increase (%)"] == 3
    assert captured["assumptions"][0]["original_percent_value"] == 0.03
    assert captured["assumptions"][0]["original_percent_format"] == "0.0%"


def test_original_quote_overlap_is_not_silently_calculated(proposal_source):
    p = Proposal()
    rows = normalise_table(
        [
            {
                "Role": "Facility Manager",
                "Quantity": 1,
                "Unit": "headcount",
                "Monthly rate": 8000,
                "Annual increase (%)": 3,
                "Year 1 rate": 8000,
                "Year 2 rate": 9000,
                "Remarks": "Weekday office coordination",
            }
        ],
        "Conflicting supplied quote",
    )
    capture_original(
        p,
        {"description": "Office coordination", "start": "2027-01-01", "base_years": 1, "option_years": 1},
        rows,
        True,
    )
    assert comparison(p, load_records())["original"] is None
    assert not p.original_quote_records
    resolve_conflict(
        p,
        p.plan.roles[0].id,
        {
            "method": "Supplied yearly prices",
            "yearly_rates": [8000, 9000],
            "source": "Officer selected the supplied yearly prices",
        },
        load_records(),
        "Officer",
    )
    assert comparison(p, load_records())["reviewed"]["contract_total"] == 204000
    assert comparison(p, load_records())["original"] is None
    assert p.plan.roles[0].escalation == [0, 0]


def test_unknown_method_and_insufficient_reference_are_not_filled_in(proposal_source):
    p = Proposal()
    rows = normalise_table(
        [
            {
                "Role": "Technician",
                "Quantity": 1,
                "Unit": "headcount",
                "Monthly rate": 8000,
                "Annual increase (%)": 3,
                "Proposed method": "Method to be clarified",
            }
        ],
        "Readable proposal",
    )
    capture_original(
        p,
        {"description": "Weekday equipment response", "start_year": 2027, "base_years": 1, "option_years": 1},
        rows,
        True,
    )
    assert p.plan.roles[0].escalation == []
    assert comparison(p, load_records())["original"] is None
    resolve_conflict(
        p,
        p.plan.roles[0].id,
        {
            "method": "General annual increase",
            "percent": 3,
            "source": "Officer clarified the supplied assumption",
        },
        load_records(),
        "Officer",
    )
    assert p.plan.roles[0].escalation == [0, 0.03]
    p = captured_example()
    rid = next(r.id for r in p.plan.roles if r.label == "Technician")
    with pytest.raises(ValueError, match="No supported general annual reference"):
        correction_preview(p, rid, [])
    proposal = correction_preview(p, rid, [], 2)
    assert not proposal["reference_available"]
    accept_correction(p, proposal, [], "Officer")
    assert next(r for r in p.plan.roles if r.id == rid).escalation == [0, 0.02]


def test_equivalent_track1_and_track2_inputs_return_identical_results(proposal_source):
    from services.cost_routing import pin_route
    from services.future_groups import general_proposal
    from services.models import Plan, Role

    p = complete_example(captured_example(), load_records())
    plan = Plan(
        profile=p.plan.profile.model_copy(deep=True),
        roles=[
            Role(id="track1-tech", label="Technician", quantity=4, quantity_unit="headcount", rate=5000),
            Role(
                id="track1-clean",
                label="Cleaner",
                quantity=2,
                quantity_unit="headcount",
                rate=3600,
                escalation=[0, 0],
                escalation_reason="Salary-only PWM",
                pwm=next(r.pwm.model_copy(deep=True) for r in p.plan.roles if r.pwm),
            ),
        ],
    )
    for role in plan.roles:
        pin_route(role)
    plan = general_proposal(plan, load_records(), ["track1-tech"], 10, "", "", True, "Officer")
    one, two = budget(plan, load_records()), budget(p.plan, load_records())
    assert one["annual_totals"] == two["annual_totals"]
    assert one["contract_total"] == two["contract_total"]
    assert one["base_total"] == two["base_total"] and one["option_total"] == two["option_total"]


def test_assistant_contract_reads_extracted_method_and_shared_route(proposal_source):
    from types import SimpleNamespace as NS

    from test_agent import MockResponses, call, response

    from services.agent import converse
    from services.models import Plan

    p = captured_example()
    before = p.plan.model_dump_json()
    tech = p.original_rows[0]
    responses = MockResponses(
        [
            response(
                [
                    call(
                        "check_future_cost_method",
                        {"role_title": tech["Role"], "requested_method": "PWM salary progression"},
                    )
                ]
            ),
            response(
                text="Technician is marked No for PWM. Review the general annual reference and confirm the Cleaner category."
            ),
        ]
    )
    converse(
        p.plan,
        load_records(),
        "Review this supplied proposal row: "
        + json.dumps({"role": tech["Role"], "method": tech["Proposed method"]}),
        NS(responses=responses),
        "mock-only",
        "Validation",
        review_existing_proposal=True,
    )
    checked = p.plan.tool_audit[0]["result"]
    assert checked["conflict"] and checked["pwm_flag"] == "No" and not checked["plan_changed"]
    assert checked["general_reference"]["average"] == pytest.approx(0.1)
    assert p.plan.roles == Plan.model_validate_json(before).roles
    assert original_signature(p) == p.original_digest


def test_loaded_example_follows_normal_shared_ui_without_paid_calls(proposal_source, monkeypatch, tmp_path):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'example-ui.db'}")
    monkeypatch.setenv("DEPLOYMENT_MODE", "local")
    monkeypatch.setenv("OPENAI_API_KEY", "")
    monkeypatch.setenv("OPENAI_MODEL", "")
    at = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=90).run()
    original_track1 = at.session_state["plan"].model_dump_json()
    at.button(key="activity_budgeting").click().run()
    at.button(key="choose_track_2").click().run()
    assert not at.dataframe and not at.checkbox
    assert not any(t.label == "Project name" for t in at.text_input)
    assert any(h.value == "Upload your existing proposal" for h in at.header)
    assert at.button(key="proposal_extract_" + at.session_state["proposal"].id).disabled
    assert not any("PWM review demonstration" in e.label for e in at.expander)
    at.button(key="t2_load_example").click().run()
    p = at.session_state["proposal"]
    assert p.original is None and not p.plan.decisions and not p.plan.confirmations
    at.checkbox(key="proposal_basis_" + p.id).check().run()
    at.checkbox(key="proposal_interpret_" + p.id).check().run()
    at.button(key="proposal_capture_" + p.id).click().run()
    assert not at.exception
    p = at.session_state["proposal"]
    assert len(p.plan.roles) == 4
    assert all(not row.get("Proposed method") for row in p.original_rows)
    assert not any(b.label.startswith("Review correction") for b in at.button)
    at.button(key="t2_confirm_deployment_" + p.id).click().run()
    assert not at.exception
    assert at.session_state["proposal"].plan.ui_step == 3
    at.button(key="confirm_rates").click().run()
    assert not at.exception
    assert at.session_state["proposal"].plan.ui_step == 4
    p = at.session_state["proposal"]
    cleaner = next(r for r in p.plan.roles if r.label == "Cleaner")
    prefix = f"pwm_{p.plan.id}_{cleaner.id}"
    from services.pwm import catalog

    schedule = next(
        s
        for s in catalog()["schedules"]
        if s["sector"] == "cleaning"
        and "Group 1" in s["category"]
        and s["job_level"] == "General / indoor cleaner"
    )
    at.selectbox(key=prefix + "_schedule_id").set_value(schedule["id"]).run()
    at.checkbox(key=prefix + "_applicability_confirmed").check().run()
    assert at.number_input(key=prefix + "_starting_salary").value == 2300
    assert not at.checkbox(key=prefix + "_remainder_flat_confirmed").value
    at.checkbox(key=prefix + "_remainder_flat_confirmed").check().run()
    apply = next(b for b in at.button if b.label == "Apply salary progression to this role")
    apply.click().run()
    assert not at.exception
    assert not at.button(key="confirm_escalation").disabled, [
        x.value
        for x in [*at.info, *at.caption]
        if "confirm" in x.value.lower() or "unresolved" in x.value.lower()
    ]
    at.button(key="confirm_escalation").click().run()
    assert not at.exception
    assert at.session_state["proposal"].stage == 2
    assert [t.label for t in at.tabs] == [
        "Budget",
        "Scenarios",
        "Manpower",
        "Assumptions & sources",
        "Changes from your original proposal",
    ]
    assert any(m.label == "Full contract manpower budget (SGD)" for m in at.metric)
    assert all(not row.get("Proposed method") for row in at.session_state["proposal"].original_rows)
    assert at.session_state["plan"].model_dump_json() == original_track1
    assert not at.session_state["proposal"].plan.tool_audit
    assert not at.button(key="proposal_final_" + p.id).disabled

    at.button(key="proposal_final_" + p.id).click().run()
    assert not at.exception and "final" in at.session_state["proposal"].plan.confirmations
