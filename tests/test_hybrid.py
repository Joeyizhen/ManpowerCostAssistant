"""Input-table, cohort mean and export regressions; synthetic records only."""

from io import BytesIO

import openpyxl
import pytest

from services.demo import demo_plan
from services.evidence import load_records, rate_benchmark
from services.export import export_excel
from services.hybrid import initialise_rates, parse_manpower_list, save_input_table, table_rows
from services.journey import working_draft
from services.models import Plan
from services.role_reference import role_titles
from services.workflow import missing_for_gate


def test_paste_preserves_combined_titles_fractional_allocations_and_custom_titles():
    rows = parse_manpower_list(
        '"Senior Property Executive cum FSM", 0.5, shared supervision\nCustom Events Support: 1\nTechnician\t4\tweekday fault attendance'
    )
    p = Plan()
    save_input_table(p, "Officer FM brief", 2028, 2, 2, rows, "Owner-appointed trade contractors")
    assert [r["Role title"] for r in table_rows(p)] == [
        "Senior Property Executive cum FSM",
        "Custom Events Support",
        "Technician",
    ]
    assert [r["Quantity / FTE"] for r in table_rows(p)] == [0.5, 1, 4]
    assert not p.roles and not p.confirmations and not p.pending
    assert p.input_support == "Owner-appointed trade contractors"
    assert working_draft(p).profile.start is None
    assert working_draft(p).roles[0].mapped_role is None
    assert "Senior Property Executive Cum FSM" in role_titles()


@pytest.mark.parametrize("text", ["Technician, -1", "Technician, nan", "Technician", "Technician, 0", ", 4"])
def test_invalid_paste_never_invents_quantities(text):
    with pytest.raises(ValueError):
        parse_manpower_list(text)


def test_dynamic_table_missing_optional_cells_are_unknown_and_visible_to_agent():
    from services.agent import execute_tool

    p = Plan()
    rows = [
        {
            "Role title": "Custom role",
            "Quantity / FTE": 2,
            "Duties / coverage notes": float("nan"),
            "Deployment": float("nan"),
        }
    ]
    save_input_table(p, "Known project brief", 2028, 2, 0, rows, "Trade contractors retained")
    r = working_draft(p).roles[0]
    assert r.duties_notes == "" and r.deployment == "unspecified"
    context = execute_tool(p, load_records(), "get_project", "{}")
    assert context["working_profile"]["description"] == "Known project brief"
    assert context["working_roles"][0]["label"] == "Custom role"
    assert context["input_support"] == "Trade contractors retained"
    assert context["officer_manpower_input"][0]["quantity"] == 2
    assert not context["roles"] and not p.confirmations


def test_average_is_arithmetic_mean_of_distinct_comparable_same_year_records():
    source = rate_benchmark(load_records(), "Technician", 2028)["records"][:3]
    rows = [{**r, "unit_rate_monthly": v} for r, v in zip(source, [100, 200, 900], strict=True)]
    rows.append(rows[0].copy())
    rows.append({**rows[1], "record_id": "foreign", "schedule_id": "foreign", "currency": "USD"})
    bm = rate_benchmark(rows, "Technician", 2028)
    assert bm["sample_count"] == 3 and bm["minimum"] == 100 and bm["maximum"] == 900
    assert bm["average"] == 400 and bm["median"] == 200
    assert not rate_benchmark(rows[:2], "Technician", 2028)["available"]
    assert rate_benchmark(rows[:2], "Technician", 2028)["average"] is None


def test_mixed_calendar_years_cannot_supply_a_rate_average():
    rows = rate_benchmark(load_records(), "Technician", 2024)["records"][:3]
    rows = [r.copy() for r in rows]
    rows[-1]["pricing_year"] = 2025
    bm = rate_benchmark(rows, "Technician")
    assert not bm["available"] and bm["average"] is None
    assert "different priced years" in bm["limitations"]


def test_defaults_are_editable_average_not_unsupported_rate_and_no_approval():
    p, records = demo_plan(), load_records()
    for r in p.roles:
        r.rate, r.rate_reason = None, ""
    initialise_rates(p, records, 2028)
    for r in p.roles[:2]:
        assert r.rate == round(
            rate_benchmark(records, r.mapped_role, 2028, r.benchmark_status, r.benchmark_scope)["average"], 2
        )
    assert p.roles[-1].rate is None and not p.confirmations
    p.roles[-1].rate = 6500
    p.roles[-1].rate_reason = ""
    assert missing_for_gate(p, "rates", records)
    from services.decision_basis import record_basis

    record_basis(
        p,
        "rate",
        [r.id for r in p.roles if r.included],
        "Quotation",
        "Synthetic quotation covers the same inclusive manpower duties at the contract start year",
        "Synthetic quotation Q1",
        "officer assumption",
        "Synthetic officer",
    )
    assert not missing_for_gate(p, "rates", records)


def test_row_removal_is_explicit_exclusion_and_preserves_other_answers():
    p = demo_plan()
    rows = table_rows(p)
    rows[0]["Quantity / FTE"] = 2
    save_input_table(p, p.profile.description, 2024, 4, 2, rows[:2])
    draft = working_draft(p)
    assert draft.roles[0].quantity == 2 and draft.roles[1].quantity == 5
    assert not draft.roles[2].included and "removed" in draft.roles[2].exclusion_reason
    assert draft.responsibilities == p.responsibilities


def test_excel_contains_arithmetic_average_support_and_optional_allowance_notes():
    p, records = demo_plan(), load_records()
    p.input_support = "Existing specialist maintenance contractor"
    p.roles[0].rate_notes = "Extra allowance for reporting workload"
    wb = openpyxl.load_workbook(BytesIO(export_excel(p, records)), data_only=True)
    headers = [c.value for c in wb["Rate evidence"][4]]
    index = headers.index("Average (arithmetic)") + 1
    assert (
        wb["Rate evidence"].cell(5, index).value
        == rate_benchmark(records, p.roles[0].mapped_role, 2024, "awarded", "Public venue FM")["average"]
    )
    assert "Extra allowance" in wb["Rate evidence"].cell(5, 4).value
    assert any(
        "Existing specialist maintenance" in str(c.value)
        for row in wb["Operational requirements"]
        for c in row
    )
    assert len(wb["Budget summary"]._charts) == 2
