"""Tests for Track 4 historical tender benchmarking."""

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from services.agent2.track3 import Track3InputError
from services.agent2.track4 import (
    MIN_COMPARABLE_RECORDS,
    SYNTHETIC_LABEL,
    adjustment_factor,
    compare_with_history,
    cost_index_snapshot,
    export_track4_excel,
    generate_review_narrative,
    historical_input_fingerprint,
    load_synthetic_history,
    synthetic_track4_tenderers,
)


def test_synthetic_history_is_labelled_and_structured_for_track5():
    history = load_synthetic_history()
    assert history["Synthetic"].all()
    assert set(history["Source Type"]) == {"synthetic historical tender record"}
    assert history["Record ID"].is_unique
    assert history["Standard Role"].nunique() >= 10
    assert {"Facility Type", "Service Category", "GFA Band", "Pricing Year"} <= set(history.columns)
    assert "NOT VERIFIED MARKET EVIDENCE" in SYNTHETIC_LABEL


def test_historical_comparison_produces_high_medium_within_and_insufficient_examples():
    result = compare_with_history(
        synthetic_track4_tenderers(), load_synthetic_history(), pricing_year_from=2022
    )
    comparisons = result.comparisons
    assert "High" in set(comparisons["Priority"])
    assert "Medium" in set(comparisons["Priority"])
    assert "Low" in set(comparisons["Priority"])
    assert "Manual review" in set(comparisons["Priority"])
    insufficient = comparisons[comparisons["Priority"] == "Manual review"].iloc[0]
    assert insufficient["Similar Historical Rates Found"] < MIN_COMPARABLE_RECORDS
    assert insufficient["Historical Comparison"] == "Insufficient similar historical data"


def test_role_facility_service_and_year_filters_are_traceable():
    result = compare_with_history(
        synthetic_track4_tenderers(),
        load_synthetic_history(),
        facility_type="Commercial",
        service_category="Integrated Facilities Management",
        gfa_band="Below 25,000 sqm",
        pricing_year_from=2024,
    )
    reasons = " ".join(result.comparisons["Why These Historical Records Were Selected"])
    assert "facility type: Commercial" in reasons
    assert "service category: Integrated Facilities Management" in reasons
    assert "pricing year: 2024 onwards" in reasons
    assert result.comparisons["Source Reference"].str.contains("SYN-HIST-").any()


def test_median_quartiles_and_differences_are_calculated_from_selected_history():
    history = load_synthetic_history()
    result = compare_with_history(synthetic_track4_tenderers(), history)
    row = result.comparisons[result.comparisons["Standard Role Used for Comparison"] == "Technician"].iloc[0]
    selected = history[(history["Standard Role"] == "Technician") & (history["Quantity Unit"] == "headcount")]
    assert row["Historical Median Monthly Rate (SGD)"] == pytest.approx(
        selected["Monthly Rate (SGD)"].median()
    )
    assert row["Typical Historical Range – Low (SGD)"] == pytest.approx(
        selected["Monthly Rate (SGD)"].quantile(0.25)
    )
    assert row["Difference from Historical Median (SGD)"] == pytest.approx(
        row["Tendered Monthly Rate (SGD)"] - row["Historical Median Monthly Rate (SGD)"]
    )


def test_quantity_unit_mismatch_returns_insufficient_evidence_not_false_comparison():
    tenders = synthetic_track4_tenderers()
    tenders["Synthetic Tenderer Alpha"] = tenders["Synthetic Tenderer Alpha"].copy()
    tenders["Synthetic Tenderer Alpha"]["Quantity Unit"] = "fte"
    result = compare_with_history(tenders, load_synthetic_history())
    assert set(result.comparisons["Priority"]) == {"Manual review"}
    assert (result.comparisons["Similar Historical Rates Found"] == 0).all()


def test_multiple_tenderers_and_questions_are_numbered_and_prioritised():
    tenders = synthetic_track4_tenderers()
    tenders["Synthetic Tenderer Beta"] = tenders["Synthetic Tenderer Alpha"].copy()
    result = compare_with_history(tenders, load_synthetic_history())
    assert set(result.summary["Tenderer"]) == set(tenders)
    assert list(result.questions["S/N"]) == list(range(1, len(result.questions) + 1))
    assert {"Priority", "Supporting Evidence", "Source Reference"} <= set(result.questions.columns)


def test_missing_or_invalid_submission_is_blocked():
    invalid = pd.DataFrame(
        [
            {
                "Role": "Technician",
                "Contract Year": 1,
                "Quantity": None,
                "Quantity Unit": "headcount",
                "Monthly Rate (SGD)": 4000,
                "Source": "test",
            }
        ]
    )
    with pytest.raises(Track3InputError, match="invalid values"):
        compare_with_history({"Tenderer": invalid}, load_synthetic_history())


def test_fingerprint_changes_when_filters_or_verified_inputs_change():
    tenders = synthetic_track4_tenderers()
    filters = {"Facility Type": "Any"}
    first = historical_input_fingerprint(tenders, None, filters)
    changed_filters = historical_input_fingerprint(tenders, None, {"Facility Type": "Commercial"})
    changed_tenders = {name: frame.copy() for name, frame in tenders.items()}
    changed_tenders["Synthetic Tenderer Alpha"].loc[0, "Quantity"] += 1
    assert first != changed_filters
    assert first != historical_input_fingerprint(changed_tenders, None, filters)


def test_export_contains_traceability_and_no_award_recommendation():
    history = load_synthetic_history()
    tenders = synthetic_track4_tenderers()
    result = compare_with_history(tenders, history)
    mappings = pd.DataFrame(
        [
            {
                "Submitted by": "Synthetic Tenderer Alpha",
                "Role title in submission": "Technician",
                "Compare as": "Technician",
            }
        ]
    )
    binary = export_track4_excel({"name": "Test"}, result, history, mappings, {"Facility Type": "Any"})
    workbook = load_workbook(BytesIO(binary), read_only=True)
    assert {
        "Historical Comparison",
        "Historical Source Records",
        "Clarification Questions",
        "Limitations",
    } <= set(workbook.sheetnames)
    limitation_text = " ".join(
        str(cell.value) for row in workbook["Limitations"] for cell in row if cell.value
    )
    assert "do not recommend an award" in limitation_text


def test_historical_rates_are_adjusted_to_target_year_without_mutating_source():
    history = load_synthetic_history()
    original = history.copy(deep=True)
    result = compare_with_history(synthetic_track4_tenderers(), history, target_pricing_year=2026)
    factor, explanation = adjustment_factor(2022, 2026, "Facilities Manager")
    assert factor > 1
    assert "2022→2026" in explanation
    assert (result.adjustments["Target Pricing Year"] == 2026).all()
    pd.testing.assert_frame_equal(history, original)


def test_wage_index_distinguishes_verified_years_from_future_scenarios():
    index = cost_index_snapshot().set_index("Year")
    assert index.loc[2023, "Status"] == "Verified"
    assert index.loc[2025, "Status"] == "Verified"
    assert index.loc[2026, "Status"] == "Scenario"
    assert index.loc[2026, "Cumulative Index (2022=1.000)"] > index.loc[2025, "Cumulative Index (2022=1.000)"]
    _, explanation = adjustment_factor(2022, 2026, "Facilities Manager")
    assert "Includes officer-approved scenario" in explanation
    assert "CPF is already within this wage-change series" in explanation


def test_pwm_floor_is_a_reference_not_an_automatic_multiplier():
    result = compare_with_history(synthetic_track4_tenderers(), load_synthetic_history())
    row = result.comparisons[result.comparisons["Standard Role"] == "Facilities Manager"].iloc[0]
    assert pd.isna(row["Applicable Statutory Wage Floor (SGD)"])
    cleaning_factor, _ = adjustment_factor(2022, 2026, "Cleaning Supervisor")
    general_factor, _ = adjustment_factor(2022, 2026, "Facilities Manager")
    assert cleaning_factor == general_factor


def test_demonstration_contains_expected_signed_role_variances_and_overall_review():
    result = compare_with_history(synthetic_track4_tenderers(), load_synthetic_history())
    by_role = result.comparisons.set_index("Standard Role")
    assert by_role.loc["Facilities Manager", "Variance to Adjusted Benchmark (%)"] == pytest.approx(
        0.06, abs=0.002
    )
    assert by_role.loc["Technical Executive", "Variance to Adjusted Benchmark (%)"] == pytest.approx(
        -0.15, abs=0.002
    )
    assert by_role.loc["Workplace / Events Manager", "Priority"] == "Manual review"
    assert not result.overall.empty
    assert "Overall Interpretation" in result.overall


def test_policy_evidence_is_dated_and_uses_official_direct_urls():
    result = compare_with_history(synthetic_track4_tenderers(), load_synthetic_history())
    assert result.policy_evidence["Date Accessed"].notna().all()
    assert result.policy_evidence["Direct URL"].str.startswith("https://").all()
    assert set(result.policy_evidence["Evidence Type"]) == {"Verified external evidence"}


def test_ai_narrative_falls_back_when_unavailable_or_invalid():
    result = compare_with_history(synthetic_track4_tenderers(), load_synthetic_history())
    assert generate_review_narrative(result).iloc[0]["Status"] == "Deterministic fallback"

    class InvalidResponses:
        def create(self, **_kwargs):
            return type("Response", (), {"output_text": "not json"})()

    client = type("Client", (), {"responses": InvalidResponses()})()
    assert (
        generate_review_narrative(result, client, "test-model").iloc[0]["Status"] == "Deterministic fallback"
    )


def test_low_rate_flags_possible_missing_cost_components():
    result = compare_with_history(synthetic_track4_tenderers(), load_synthetic_history())
    row = result.comparisons[result.comparisons["Standard Role"] == "Technical Executive"].iloc[0]
    assert "Possible omitted or bundled components" in row["What Remains Uncertain"]
