"""Repeatable workbook ingestion. The original workbook is opened read-only."""

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import UTC, date, datetime
from pathlib import Path

from services.reference_data import annual_reference, mapping_safe, root

IMPORT_VERSION = "sprint2-reference-v2"

REQUIRED_COLUMNS = {
    "Manpower Data": {
        "Record ID",
        "Tender Reference",
        "Tenderer",
        "Property Name",
        "Original Role Title",
        "Unit Rate per Month (SGD)",
        "Quantity of Manpower",
        "Quoted Period / Contract Year",
    },
    "Record Controls": {
        "Record ID",
        "Mapping ID",
        "Role Line ID",
        "Sanitised Source Reference",
        "Completed Rate Basis",
        "Completed Quantity Basis",
        "Completed Rate Eligibility",
        "Completed Quantity Eligibility",
    },
    "Standard Roles": {"Role Title", "Typical manpower?", "PWM applies?", "Scope"},
    "Role Mapping Review": {
        "Mapping ID",
        "Original Role Title",
        "Proposed Standard Role",
        "Mapping Review",
        "Combined Duties",
        "Source Section",
    },
    "Tender Details Data Base": {
        "Tender Reference",
        "Tenderer",
        "Property Name",
        "Tender ID",
        "Submission ID",
        "Schedule ID",
        "Property ID",
    },
    "Source Register": {"Public Document Label", "Submission ID", "Source ID"},
}


def structural_checks(data):
    errors = []
    for name, columns in REQUIRED_COLUMNS.items():
        absent = columns - set(data[name][0] if data[name] else {})
        if absent:
            errors.append(f"{name}: missing columns {', '.join(sorted(absent))}")
    roles = data["Standard Roles"]
    titles = [str(r.get("Role Title") or "").strip().casefold() for r in roles]
    if not all(titles) or len(titles) != len(set(titles)):
        errors.append("Standard Roles: blank or duplicate role titles")
    for r in roles:
        for field in ("Typical manpower?", "PWM applies?"):
            if str(r.get(field) or "").strip().casefold() not in ("yes", "no"):
                errors.append(f"{r.get('Role Title')}: missing/invalid {field}")
    return errors


def validate_activation(snapshot):
    report = snapshot["readiness_report"]
    if report.get("structural_errors") or report["identifier_problems"] or report["join_problems"]:
        raise ValueError("Snapshot has structural, identifier or relationship errors; activation refused")
    if not snapshot["records"]:
        raise ValueError("Cannot activate an empty reference snapshot")


def serial(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def table(sheet, header_row=1):
    iterator = sheet.iter_rows(min_row=header_row, values_only=True)
    headers = next(iterator)
    return [
        {str(k): serial(v) for k, v in zip(headers, row, strict=False) if k is not None}
        for row in iterator
        if any(v is not None for v in row)
    ]


def unique_index(rows, key, issues):
    counts = Counter(r.get(key) for r in rows)
    invalid = {k for k, n in counts.items() if not k or n != 1}
    if invalid:
        issues.append({"table_key": key, "invalid_or_duplicate_keys": sorted(map(str, invalid))})
    return {r[key]: r for r in rows if r.get(key) not in invalid}


def positive(value):
    try:
        result = float(value)
        return result > 0 and result < float("inf")
    except (TypeError, ValueError):
        return False


def pwm_review(rows):
    from services.pwm import catalog

    existing = catalog()["schedules"]
    matched, unresolved = [], []
    for row in rows:
        candidates = [
            s
            for s in existing
            if s["source"] == row.get("Official Source URL")
            and s["job_level"].casefold() == str(row.get("Job Level", "")).casefold()
            and row.get("Wage Basis") == "Basic"
            and row.get("Unit") == "SGD/month"
            and any(
                p["effective"] == str(row.get("Effective From", ""))[:10]
                and p["wage"] == row.get("Minimum Wage SGD")
                for p in s["points"]
            )
        ]
        item = {
            "schedule_id": row.get("Schedule ID"),
            "series_id": row.get("Series ID"),
            "job_level": row.get("Job Level"),
            "basis": row.get("Wage Basis"),
            "unit": row.get("Unit"),
            "effective": row.get("Effective From"),
            "source": row.get("Official Source URL"),
        }
        if len(candidates) == 1:
            matched.append(
                {
                    **item,
                    "existing_catalogue_id": candidates[0]["id"],
                    "note": "Matching existing wage observation retained; category/applicability still requires officer confirmation",
                }
            )
        else:
            unresolved.append(
                {
                    **item,
                    "reason": "No unique exact existing job/source/monthly-basic wage observation; catalogue not overwritten",
                }
            )
    return {
        "observations": len(rows),
        "matching_observations": matched,
        "differences_requiring_review": unresolved,
        "catalogue_updated": False,
        "note": "Workbook policy observations remain separate. Its Date Checked is supplied metadata, not a new verification by this importer.",
    }


def build_snapshot(workbook):
    from openpyxl import load_workbook

    path = Path(workbook)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    book = load_workbook(path, read_only=True, data_only=True)
    required = [
        "Manpower Data",
        "Tender Details Data Base",
        "Standard Roles",
        "Record Controls",
        "Role Mapping Review",
        "Source Register",
        "Ambiguous Keys",
        "Issue Examples",
        "Readiness and Issues",
        "Data Dictionary",
        "PWM Trend",
    ]
    if any(name not in book.sheetnames for name in required):
        raise ValueError("The prepared workbook is missing required relationship/control sheets")
    data = {name: table(book[name]) for name in required if name != "PWM Trend"}
    pwm_header = next(
        (i for i, row in enumerate(book["PWM Trend"].iter_rows(values_only=True), 1) if "Schedule ID" in row),
        None,
    )
    if pwm_header is None:
        raise ValueError("PWM Trend has no schedule table")
    pwm = table(book["PWM Trend"], pwm_header)
    book.close()
    # Formatted trailing spreadsheet rows never become role choices.
    data["Standard Roles"] = [r for r in data["Standard Roles"] if str(r.get("Role Title") or "").strip()]
    structure = structural_checks(data)
    issues = []
    controls = unique_index(data["Record Controls"], "Record ID", issues)
    mappings = unique_index(data["Role Mapping Review"], "Mapping ID", issues)
    unique_records = unique_index(data["Manpower Data"], "Record ID", issues)
    sources = unique_index(data["Source Register"], "Public Document Label", issues)
    detail_keys = [
        tuple(r.get(k) for k in ("Tender ID", "Submission ID", "Schedule ID", "Property ID"))
        for r in data["Tender Details Data Base"]
    ]
    if any(not all(key) for key in detail_keys) or len(detail_keys) != len(set(detail_keys)):
        structure.append("Tender details contain blank IDs or duplicate relationship keys")
    if set(controls) != set(unique_records):
        structure.append("Record Controls and Manpower Data Record IDs do not match")
    details = defaultdict(list)
    for r in data["Tender Details Data Base"]:
        details[(r.get("Tender Reference"), r.get("Tenderer"), r.get("Property Name"))].append(r)
    taxonomy = {r["Role Title"]: r for r in data["Standard Roles"]}
    if any(
        m.get("Proposed Standard Role") and m["Proposed Standard Role"] not in taxonomy
        for m in mappings.values()
    ):
        structure.append("Role Mapping Review contains an unknown Standard Role")
    dataset_id = digest[:16] + "-" + IMPORT_VERSION
    source_positions = {r.get("Record ID"): i for i, r in enumerate(data["Manpower Data"], 2)}
    records, missing = [], Counter()
    for record_id, raw in unique_records.items():
        c = controls.get(record_id, {})
        m = mappings.get(c.get("Mapping ID"), {})
        candidates = details.get(
            (raw.get("Tender Reference"), raw.get("Tenderer"), raw.get("Property Name")), []
        )
        d = candidates[0] if len(candidates) == 1 else {}
        public_source = str(c.get("Sanitised Source Reference") or "").split(" | ")[0]
        source = sources.get(public_source, {})
        joined = bool(c and m and d and source and source.get("Submission ID") == d.get("Submission ID"))
        if not joined:
            missing["unresolved_record_relationship"] += 1
        year_match = re.fullmatch(r"Year\s+(\d+)", str(raw.get("Quoted Period / Contract Year", "")), re.I)
        year = int(year_match[1]) if year_match else None
        flags = str(c.get("Issue Flags") or "")
        safe_title = mapping_safe(m)
        rate = raw.get("Unit Rate per Month (SGD)")
        quantity = raw.get("Quantity of Manpower")
        rate_origin = c.get("Completed Rate Basis", "Unknown")
        quantity_origin = c.get("Completed Quantity Basis", "Unknown")
        rate_ok = (
            joined
            and safe_title
            and year is not None
            and positive(rate)
            and rate_origin in ("Source-extracted", "Source-derived")
            and c.get("Completed Rate Eligibility") == "Conditional source-rate candidate"
        )
        qty_ok = (
            joined
            and safe_title
            and year is not None
            and positive(quantity)
            and quantity_origin in ("Source-extracted", "Source-derived")
            and c.get("Completed Quantity Eligibility") == "Conditional source-quantity candidate"
            and float(quantity).is_integer()
        )
        standard = m.get("Proposed Standard Role")
        records.append(
            {
                "record_id": record_id,
                "record_revision": 1,
                "reference_kind": "prepared quotation",
                "dataset_id": dataset_id,
                "tender_id": d.get("Tender ID"),
                "submission_id": d.get("Submission ID"),
                "schedule_id": d.get("Schedule ID"),
                "property_id": d.get("Property ID"),
                "role_line_id": c.get("Role Line ID"),
                "mapping_id": c.get("Mapping ID"),
                "source_id": source.get("Source ID"),
                "source_sha256": source.get("SHA256"),
                "raw_role_name": raw.get("Original Role Title"),
                "standard_role": standard,
                "source_section": m.get("Source Section"),
                "role_scope": taxonomy.get(standard, {}).get("Scope"),
                "mapping_review": m.get("Mapping Review"),
                "combined_duties": m.get("Combined Duties"),
                "quantity": quantity,
                "quantity_unit": "Source count/allocation; headcount/FTE unconfirmed",
                "quantity_origin": quantity_origin,
                "unit_rate_monthly": rate,
                "rate_origin": rate_origin,
                "currency": "SGD",
                "rate_basis": "monthly contractor unit quote; source unit unconfirmed",
                "contract_year": year,
                "contract_duration": None,
                "pricing_year": None,
                "submission_year": None,
                "contract_start_year": None,
                "calendar_date_basis": "Synthetic",
                "completed_period_from": raw.get("Period From"),
                "completed_period_to": raw.get("Period To"),
                "award_status": d.get("Submission Award Status") or "Unknown",
                "scope": None,
                "inclusions": "Unknown",
                "exclusions": "Unknown",
                "validation_status": "conditional prepared reference",
                "data_origin": "prepared source quotation; field origins retained",
                "source": c.get("Sanitised Source Reference"),
                "source_file": public_source,
                "source_row": f"Manpower Data row {source_positions[record_id]}; CSV row {c.get('Full CSV Row')}",
                "experience": None,
                "qualifications": None,
                "issue_flags": flags,
                "eligibility": {
                    "rate": bool(rate_ok),
                    "annual": bool(rate_ok and "Partial period" not in flags),
                    "quantity": bool(qty_ok),
                },
                "original": raw,
                "record_control": c,
                "mapping": m,
            }
        )
    durations = defaultdict(int)
    for r in records:
        if r["schedule_id"] and r["contract_year"]:
            durations[r["schedule_id"]] = max(durations[r["schedule_id"]], r["contract_year"])
    for r in records:
        r["contract_duration"] = durations.get(r["schedule_id"])
    annual = annual_reference(records, general=True)
    report = {
        "source_sha256": digest,
        "dataset_id": dataset_id,
        "imported_records": len(records),
        "sheet_records": {name: len(rows) for name, rows in data.items()},
        "identifier_problems": issues,
        "structural_errors": structure,
        "join_problems": dict(missing),
        "detail_join": "Record ID to controls; unique tender/tenderer/property aliases resolve supplied detail IDs; source-register Submission ID cross-check. Ambiguous joins are ineligible, never first-match guesses.",
        "distinct_tenders": len({r["tender_id"] for r in records if r["tender_id"]}),
        "distinct_submissions": len({r["submission_id"] for r in records if r["submission_id"]}),
        "distinct_schedules": len({r["schedule_id"] for r in records if r["schedule_id"]}),
        "distinct_role_lines": len({r["role_line_id"] for r in records if r["role_line_id"]}),
        "eligible_by_use": {
            use: sum(r["eligibility"][use] for r in records) for use in ("rate", "annual", "quantity")
        },
        "excluded_synthetic_rates": sum(r["rate_origin"] == "Synthetic" for r in records),
        "synthetic_quantities": sum(r["quantity_origin"] == "Synthetic" for r in records),
        "unresolved_or_combined_mapping_records": sum(not mapping_safe(r["mapping"]) for r in records),
        "general_annual_pairs": annual["pair_count"],
        "general_annual_submissions": annual["submission_count"],
        "general_annual_lines": len(annual["schedules"]),
        "general_annual_average": annual["average"],
        "calendar_dates": "All completed dates synthetic; actual pricing years unavailable",
        "outcomes": dict(Counter(r["award_status"] for r in records)),
        "pwm_review": pwm_review(pwm),
        "note": "Only this workbook selection imported; overlapping full CSVs were not appended. Eligibility is conditional prototype use, not full validation or award certification.",
    }
    return {
        "schema_version": "1",
        "importer_version": IMPORT_VERSION,
        "dataset_id": dataset_id,
        "source_name": path.name,
        "source_sha256": digest,
        "imported_at": datetime.now(UTC).isoformat(),
        "records": records,
        "standard_roles": data["Standard Roles"],
        "mappings": data["Role Mapping Review"],
        "sources": data["Source Register"],
        "pwm_observations": pwm,
        "readiness_report": report,
        "source_notes": {
            name: data[name]
            for name in ("Ambiguous Keys", "Issue Examples", "Readiness and Issues", "Data Dictionary")
        },
    }


def activate(dataset_id, directory=None):
    directory = Path(directory) if directory else root()
    if dataset_id != "synthetic":
        target = directory / "snapshots" / (dataset_id + ".json")
        if target.resolve().parent != (directory / "snapshots").resolve() or not target.exists():
            raise ValueError("Unknown snapshot")
        validate_activation(json.loads(target.read_text(encoding="utf-8")))
    pointer = directory / "active.json"
    old = json.loads(pointer.read_text(encoding="utf-8")) if pointer.exists() else {"active": "synthetic"}
    if old["active"] == dataset_id:
        return old
    directory.mkdir(parents=True, exist_ok=True)
    value = {"active": dataset_id, "previous": old["active"], "activated_at": datetime.now(UTC).isoformat()}
    temp = pointer.with_suffix(".tmp")
    temp.write_text(json.dumps(value, indent=2), encoding="utf-8")
    temp.replace(pointer)
    return value


def import_workbook(workbook, directory=None, activate_now=False):
    directory = Path(directory) if directory else root()
    snapshot = build_snapshot(workbook)
    target = directory / "snapshots" / (snapshot["dataset_id"] + ".json")
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        previous = json.loads(target.read_text(encoding="utf-8"))
        if (
            previous["source_sha256"] != snapshot["source_sha256"]
            or previous["importer_version"] != IMPORT_VERSION
        ):
            raise ValueError("Snapshot identifier conflict")
        snapshot = previous
    else:
        with target.open("x", encoding="utf-8") as stream:
            json.dump(snapshot, stream, ensure_ascii=False)
    (directory / (snapshot["dataset_id"] + "-readiness.json")).write_text(
        json.dumps(snapshot["readiness_report"], indent=2), encoding="utf-8"
    )
    if activate_now:
        activate(snapshot["dataset_id"], directory)
    return snapshot
