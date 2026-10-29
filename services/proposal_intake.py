"""Three-sheet proposal extraction. Values are private inputs, never reference records."""

import io
import re
from datetime import date, datetime
from pathlib import Path

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "data" / "templates"
SHEETS = {"Project", "Proposed manpower", "Proposed escalation"}
PROJECT_FIELDS = {
    "project name": "name",
    "operational brief": "description",
    "anticipated start date or month": "timing",
    "base contract duration (years)": "base_years",
    "optional extension duration (years)": "option_years",
}


def template_bytes(example=False):
    return (
        TEMPLATE_DIR / ("track2-example-proposal.xlsx" if example else "track2-blank-template.xlsx")
    ).read_bytes()


def preplanned_sample_bytes():
    return (TEMPLATE_DIR / "track2-preplanned-budget-sample.xlsx").read_bytes()


def text(value):
    return "" if value is None else str(value).strip()


def numeric(value, location):
    from services.proposals import number

    try:
        return number(value)
    except (ValueError, TypeError):
        raise ValueError(f"{location}: enter a number, or leave the field blank.") from None


def timing(value):
    if not text(value):
        return {}
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return {"start": value.isoformat(), "start_year": value.year, "start_month": value.month}
    raw = text(value)
    if re.fullmatch(r"\d{4}", raw) and 1900 <= int(raw) <= 2200:
        return {"start": None, "start_year": int(raw), "start_month": None}
    try:
        return timing(date.fromisoformat(raw))
    except ValueError:
        for fmt in ("%b %Y", "%B %Y", "%Y-%m"):
            try:
                parsed = datetime.strptime(raw, fmt)
                return {"start": None, "start_year": parsed.year, "start_month": parsed.month}
            except ValueError:
                continue
    raise ValueError("Project: anticipated start must be YYYY-MM-DD or a month/year such as Jan 2027.")


def method(value):
    key = text(value).casefold()
    return {
        "pwm": "PWM",
        "pwm salary progression": "PWM",
        "general annual increase": "General annual increase",
        "general": "General annual increase",
        "supplied yearly prices": "Supplied yearly prices",
        "quoted schedule": "Supplied yearly prices",
    }.get(key, text(value))


def series(value, location):
    if not text(value):
        return None
    return [numeric(v.strip(), location) for v in re.split(r"[,;]", text(value))]


def sheet_table(sheet, first_header, required=()):
    rows = list(sheet.iter_rows(values_only=True))
    found = next(
        (
            i
            for i, row in enumerate(rows)
            if text(row[0] if row else None).casefold() == first_header.casefold()
        ),
        None,
    )
    if found is None:
        raise ValueError(f"{sheet.title}: missing the '{first_header}' header row.")
    headers = [text(v) for v in rows[found]]
    missing = set(required) - set(headers)
    if missing:
        raise ValueError(sheet.title + ": missing columns " + ", ".join(sorted(missing)))
    return [
        (i + 1, {h: row[j] if j < len(row) else None for j, h in enumerate(headers) if h})
        for i, row in enumerate(rows[found + 1 :], found + 1)
        if any(text(v) for v in row)
    ]


def extract_template(data, name):
    from openpyxl import load_workbook

    from services.proposals import normalise_table

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        present = SHEETS & set(book.sheetnames)
        if present != SHEETS:
            return None  # Other supported Excel layouts use the existing generic reader.
        profile = {}
        for rowno, row in sheet_table(book["Project"], "Field", ("Field", "Your value")):
            field = PROJECT_FIELDS.get(text(row.get("Field")).casefold())
            value = row.get("Your value")
            if field == "timing":
                profile.update(timing(value))
            elif field in ("base_years", "option_years"):
                n = numeric(value, f"Project!B{rowno}")
                if n is not None and not n.is_integer():
                    raise ValueError(f"Project!B{rowno}: enter whole years.")
                profile[field] = int(n) if n is not None else None
            elif field:
                profile[field] = text(value)
        rows = []
        for rowno, row in sheet_table(
            book["Proposed manpower"],
            "Role",
            (
                "Role",
                "Proposed quantity",
                "Unit",
                "Monthly unit rate (SGD)",
                "Operational rationale / deployment remarks",
            ),
        ):
            if not text(row.get("Role")):
                if any(text(row.get(k)) for k in ("Proposed quantity", "Monthly unit rate (SGD)")):
                    raise ValueError(f"Proposed manpower!A{rowno}: enter the role for these figures.")
                continue
            values = {
                "Role": row.get("Role"),
                "Quantity": numeric(row.get("Proposed quantity"), f"Proposed manpower!B{rowno}"),
                "Unit": row.get("Unit"),
                "Monthly rate": numeric(row.get("Monthly unit rate (SGD)"), f"Proposed manpower!D{rowno}"),
                "Remarks": row.get("Operational rationale / deployment remarks"),
                "Source location": f"{name}, Proposed manpower!A{rowno}:F{rowno}",
            }
            entry = normalise_table([values], name)[0]
            entry["Current quantity"] = numeric(
                row.get("Current quantity (optional)"), f"Proposed manpower!F{rowno}"
            )
            entry["raw"] = {k: text(v) for k, v in row.items()}
            rows.append(entry)
        assumptions = []
        for rowno, row in sheet_table(
            book["Proposed escalation"],
            "Roles or group covered",
            (
                "Roles or group covered",
                "Proposed escalation method",
                "Annual percentage (%)",
                "Yearly increases (%)",
                "Yearly monthly rates (SGD)",
                "Explanation / document reference",
            ),
        ):
            if not any(text(v) for v in row.values()):
                continue
            percentage = numeric(row.get("Annual percentage (%)"), f"Proposed escalation!C{rowno}")
            percent_format = book["Proposed escalation"].cell(rowno, 3).number_format
            if percentage is not None and "%" in percent_format:
                percentage *= 100
            assumptions.append(
                {
                    "target": text(row.get("Roles or group covered")),
                    "method": method(row.get("Proposed escalation method")),
                    "percent": percentage,
                    "original_percent_value": row.get("Annual percentage (%)"),
                    "original_percent_format": percent_format,
                    "yearly_increases": series(
                        row.get("Yearly increases (%)"), f"Proposed escalation!D{rowno}"
                    ),
                    "yearly_rates": series(
                        row.get("Yearly monthly rates (SGD)"), f"Proposed escalation!E{rowno}"
                    ),
                    "note": text(row.get("Explanation / document reference")),
                    "source": f"{name}, Proposed escalation!A{rowno}:F{rowno}",
                }
            )
        unmatched = attach_assumptions(rows, assumptions)
        return {
            "profile": profile,
            "rows": rows,
            "assumptions": assumptions,
            "unmatched_assumptions": unmatched,
            "extraction_note": "Stored Excel values read without running macros or recalculating formulas. Confirm the captured proposal; blank values remain unanswered.",
        }
    finally:
        book.close()


def attach_assumptions(rows, assumptions):
    from services.cost_routing import resolve_route
    from services.models import Role
    from services.role_catalogue import identity

    unmatched = []
    for assumption in assumptions:
        target = assumption["target"].casefold()
        chosen = []
        for row in rows:
            scope = resolve_route(Role(id=row["id"], label=row["Role"]))
            labels = [text(v) for v in re.split(r"[;|]", assumption["target"])]
            if (
                target in ("all", "all roles")
                or (
                    target in ("all non-pwm roles", "general manpower", "non-pwm roles")
                    and scope["route"] == "general"
                )
                or any(identity(row["Role"]) == identity(label) for label in labels)
            ):
                row.setdefault("Supplied assumptions", []).append(dict(assumption))
                chosen.append(row)
        if not chosen:
            unmatched.append(assumption)
    for row in rows:
        supplied = row.get("Supplied assumptions", [])
        distinct = {
            (
                a["method"],
                a["percent"],
                tuple(a.get("yearly_increases") or []),
                tuple(a.get("yearly_rates") or []),
            )
            for a in supplied
        }
        if len(distinct) > 1:
            row["Assumption conflict"] = True
        elif supplied:
            a = supplied[0]
            row["Proposed method"] = a["method"]
            row["Annual increase (%)"] = a["percent"]
            row["Yearly rates"] = a["yearly_rates"]
            row["Yearly increases"] = a["yearly_increases"]
    return unmatched


def load_example():
    """Retained error-case fixture for regression checks, never the normal UI sample."""
    return _load_proposal_sample("track2-example-proposal.xlsx", template_bytes(True))


def load_preplanned_sample():
    return _load_proposal_sample("track2-preplanned-budget-sample.xlsx", preplanned_sample_bytes())


def _load_proposal_sample(name, data):
    from services.proposals import Proposal, extract_document

    p = Proposal()
    doc = extract_document(name, data)
    p.documents, p.input_rows, p.input_profile = [doc], doc["rows"], doc["profile"]
    p.name = doc["profile"]["name"]
    p.drafts["original_editor_seed"] = p.input_rows
    p.drafts["proposal_context_" + p.id] = p.input_profile
    p.drafts["example"] = True
    return p
