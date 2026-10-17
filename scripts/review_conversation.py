"""Reconcile a saved completed conversation locally. Never calls OpenAI."""

import argparse
import json
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from pathlib import Path

import openpyxl

from services.calculations import budget, sensitivity
from services.evidence import escalation_benchmark, load_records, rate_benchmark
from services.export import export_excel
from services.models import Plan


def review(trace):
    entries = [json.loads(line) for line in trace.read_text(encoding="utf-8").splitlines()]
    saved = next(e for e in reversed(entries) if e["kind"] == "completed")
    plan, records = Plan.model_validate(saved["plan_snapshot"]), load_records()
    result = budget(plan, records)
    assert {k: result[k] for k in saved["budget"]} == saved["budget"], (
        "Current budget differs from saved live result"
    )
    assert "final" in plan.confirmations and saved["officer_final_confirmation"]
    assert len(result["annual_totals"]) == plan.profile.years
    assert round(sum(result["annual_totals"]), 2) == result["contract_total"]
    assert round(result["base_total"] + result["option_total"], 2) == result["contract_total"]
    assert (
        round(result["waterfall"]["baseline"] + sum(result["waterfall"]["additions"]), 2)
        == result["contract_total"]
    )
    # Independent arithmetic cross-check for this saved assumption-only, fixed-quantity case.
    for role in (r for r in plan.roles if r.included):
        assert not role.policies and role.escalation_mode == "assumption" and not role.yearly_quantities
        unit = Decimal(str(role.rate))
        rows = [r for r in result["rows"] if r["role_id"] == role.id]
        for i, row in enumerate(rows):
            unit = (unit * (1 + Decimal(str(role.escalation[i])))).quantize(Decimal("0.01"), ROUND_HALF_UP)
            annual = (unit * Decimal(str(role.quantity)) * 12).quantize(Decimal("0.01"), ROUND_HALF_UP)
            assert float(annual) == row["annual_cost"]
        assert rate_benchmark(
            records, role.mapped_role, role.benchmark_year, role.benchmark_status, role.benchmark_scope
        )["available"]
        assert escalation_benchmark(records, role.mapped_role, role.benchmark_status, role.benchmark_scope)[
            "available"
        ]
    scenarios = sensitivity(plan, records)
    assert scenarios[1]["contract_total"] == result["contract_total"]
    assert scenarios[0]["contract_total"] < result["contract_total"] < scenarios[2]["contract_total"]
    original_workbook = Path(saved["workbook"])
    workbook_checks = []
    for name, payload in [
        ("preserved", original_workbook.read_bytes()),
        ("current export", export_excel(plan, records)),
    ]:
        values = openpyxl.load_workbook(BytesIO(payload), data_only=True)
        formulas = openpyxl.load_workbook(BytesIO(payload), data_only=False)
        assert len(values.sheetnames) == 9
        assert len(formulas["Budget summary"]._charts) == 2
        total_row = 5 + sum(r.included for r in plan.roles)
        assert (
            values["Budget summary"].cell(total_row, plan.profile.years + 2).value == result["contract_total"]
        )
        for i, row in enumerate(result["rows"], 5):
            assert values["Annual calculations"][f"O{i}"].value == row["annual_cost"]
            assert formulas["Annual calculations"][f"O{i}"].data_type == "f"
        workbook_checks.append(name)
    report = {
        "source_trace": str(trace),
        "zero_api_calls": True,
        "original_usage": saved["counts"],
        "contract_total": result["contract_total"],
        "base_total": result["base_total"],
        "option_total": result["option_total"],
        "annual_totals": result["annual_totals"],
        "passed": [
            "saved-state equality",
            "independent role/year arithmetic",
            "supported rate/escalation retrieval",
            "base/option/full reconciliation",
            "waterfall reconciliation",
            "sensitivity ordering",
            "nine-sheet exports, cached values, formulas and two charts",
        ],
        "workbooks_checked": workbook_checks,
        "limitations": ["Native Excel rendering/recalculation not tested", "Synthetic evidence only"],
    }
    target = trace.with_name(trace.stem + "-local-review.json")
    target.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("trace", type=Path)
    review(parser.parse_args().trace)
